#!/usr/bin/env python3
"""RAT folders, band/group recommend rules, Parameter Name, and smart matching."""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import compare_reference_parameters as audit  # noqa: E402
import license_control as lic  # noqa: E402
import smart_match as smart  # noqa: E402


def _write_sheet(wb, name, rows):
    ws = wb.create_sheet(name)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            ws.cell(r, c, val)
    return ws


def _save(wb, path: Path):
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def test_rats_and_split():
    assert smart.normalize_rats("4g, nr") == ["4G", "5G"]
    assert smart.split_parameter_id("SpectralETBasedLoadEvalSw@EnhancedMlbAlgoSwitch") == (
        "EnhancedMlbAlgoSwitch",
        "SpectralETBasedLoadEvalSw",
    )
    assert smart.split_parameter_id("InterFreqIdleMlbUeNumThd") == ("InterFreqIdleMlbUeNumThd", None)
    attr, bit = smart.split_parameter_id("CaUserLoadTransferSw@EnhancedMlbAlgoSwitch")
    assert attr == "EnhancedMlbAlgoSwitch"
    assert bit == "CaUserLoadTransferSw"


def test_band_families():
    assert smart.family_for_band(8) == "L900"
    assert smart.family_for_band(8.0) == "L900"
    assert smart.family_for_band("8.0") == "L900"
    assert smart.family_for_band("Frequency Band 8") == "L900"
    assert smart.family_for_band("L9") == "L900"
    assert smart.family_for_band("L09") == "L900"
    assert smart.family_for_band("L900") == "L900"
    assert smart.family_for_band("SITE_L09A") == "L900"
    assert smart.family_for_band("3") == "L1800"
    assert smart.family_for_band("L18") == smart.family_for_band("L1800") == "L1800"
    assert smart.family_for_band(1) == "L2100"
    assert smart.family_for_band("L21") == smart.family_for_band("L2100") == "L2100"
    assert smart.family_for_band(41) == "L2600"
    assert smart.family_for_band("N41") == "L2600"
    assert smart.family_for_band("L26") == smart.family_for_band("L2600") == "L2600"
    assert smart.family_for_band(7) == "L2600"
    # Frequency Band 3 is L18/1800, never L26/2600
    assert smart.family_for_band(3) != "L2600"
    assert smart.band_matches("L09", smart.tokens_for_family("L900", 8))
    assert smart.band_matches("L9", smart.tokens_for_family("L900", 8))
    assert smart.band_matches("L900", smart.tokens_for_family("L900", "Frequency Band 8"))
    assert smart.band_matches("L26", smart.tokens_for_family("L2600", "N41"))
    assert smart.band_matches("L2600", smart.tokens_for_family("L2600", 41))
    assert not smart.band_matches("L18", smart.tokens_for_family("L900", 8))
    assert not smart.band_matches("L26", smart.tokens_for_family("L1800", 3))
    assert smart.family_for_band("DHAPT08") is None
    import time
    toks = smart.tokens_for_family("L900", 8)
    t0 = time.perf_counter()
    for _ in range(20000):
        assert smart.band_matches("L9", toks)
        assert smart.band_matches("L09", toks)
        assert smart.band_matches("L900", toks)
    elapsed = time.perf_counter() - t0
    assert elapsed < 1.0, f"band alias matching too slow: {elapsed:.2f}s"


def test_recommend_parser():
    spec = smart.parse_recommend("L9: 1\nL18: 20\nL21:20\nL26: 35")
    assert spec.has_conditional
    ctx = smart.RowContext(band_family="L1800", band_tokens=smart.tokens_for_family("L1800", 3))
    value, note = smart.select_recommend(spec, ctx)
    assert value == "20"
    assert "L18" in note

    ctx900 = smart.RowContext(band_family="L900", band_tokens=smart.tokens_for_family("L900", 8))
    value, _note = smart.select_recommend(spec, ctx900)
    assert value == "1"

    oneline = smart.parse_recommend("L9: -74 L18:-118 L21:-118 L26:-115")
    assert oneline.has_conditional
    assert len([r for r in oneline.rules if r.band_token]) == 4
    value, note = smart.select_recommend(oneline, ctx900)
    assert value == "-74"
    ctx18 = smart.RowContext(band_family="L1800", band_tokens=smart.tokens_for_family("L1800", 3))
    value, _note = smart.select_recommend(oneline, ctx18)
    assert value == "-118"
    ctx26 = smart.RowContext(band_family="L2600", band_tokens=smart.tokens_for_family("L2600", 41))
    value, _note = smart.select_recommend(oneline, ctx26)
    assert value == "-115"

    aliases = smart.parse_recommend("L900:-74 L1800:-118 L2100:-118 L2600:-115")
    assert aliases.has_conditional
    dump8 = smart.RowContext(band_family="L900", band_tokens=smart.tokens_for_family("L900", 8.0))
    value, _note = smart.select_recommend(aliases, dump8)
    assert value == "-74"

    grouped = smart.parse_recommend(
        "(InterFreqHoGroupId=1)=>L09=-108\n"
        "(InterFreqHoGroupId=1)=>L18=-108\n"
        "(InterFreqHoGroupId=1)=>640ms\n"
        "(InterRatHoCommGroupId=1)=>L09=-112\n"
        "(InterRatHoCommGroupId=1)=>L26=-110\n"
        "(InterRatHoCommGroupId=2)"
    )
    ctx = smart.RowContext(
        band_family="L900",
        band_tokens=smart.tokens_for_family("L900", 8),
        groups={"INTERFREQHOGROUPID": {"1"}, "INTERRATHOCOMMGROUPID": {"1"}},
    )
    value, note = smart.select_recommend(grouped, ctx)
    assert value == "-108"
    ctx_ms = smart.RowContext(
        band_tokens=set(),
        groups={"INTERFREQHOGROUPID": {"1"}},
    )
    value, _note = smart.select_recommend(grouped, ctx_ms)
    assert value == "640ms"
    ctx_ir = smart.RowContext(
        band_family="L2600",
        band_tokens=smart.tokens_for_family("L2600", 41),
        groups={"INTERRATHOCOMMGROUPID": {"1"}},
    )
    value, _note = smart.select_recommend(grouped, ctx_ir)
    assert value == "-110"


def test_three_identity_columns():
    headers = ["MML Object", "Parameter ID", "Parameter Name", "Golden"]
    mapping = audit.detect_ref_header_map(headers)
    assert mapping["mml"] == 0
    assert mapping["pid"] == 1
    assert mapping["param_name"] == 2
    assert mapping["recommend"] == 3
    assert audit.is_identity_ref_header(mapping)

    odd = ["MO Name", "Param ID", "Param Name", "Default Value", "Target Value"]
    mapping2 = audit.detect_ref_header_map(odd)
    assert "mml" in mapping2 and "pid" in mapping2 and "param_name" in mapping2
    assert "recommend" in mapping2
    # Default Value must not be treated as the recommend/proposed value
    assert mapping2["recommend"] == 4

    mobility = [
        "MO Name",
        "Parameter ID",
        "Parameter Name",
        "Default Value",
        "Proposed Value",
        "Conditions",
        "Targ. t/Impact",
    ]
    mapping3 = audit.detect_ref_header_map(mobility)
    assert mapping3["recommend"] == 4
    assert mapping3["conditions"] == 5
    assert mapping3["recommend"] != mapping3["conditions"]


def test_fuzzy_names():
    hit, score, why = smart.closest_name("CelAlgoSwitch", ["CellAlgoSwitch", "CellMLB"], cutoff=0.8)
    assert hit == "CellAlgoSwitch"
    hit, score, why = smart.closest_name("NRDU CEL", ["NRDUCell", "Cell"], cutoff=0.8)
    assert smart.norm_key(hit) == "NRDUCELL"


def test_a1_not_matched_as_a2():
    headers = [
        "AAAS Based Interfreq A1 RSRP Threshold(dBm)",
        "AAAS Based Interfreq A2 RSRP Threshold(dBm)",
    ]
    hit, score, why = smart.closest_name("AAAS Based Interfreq A1 RSRP Threshold", headers, cutoff=0.8)
    assert "A1" in hit
    assert "A2" not in hit
    assert why == "exact"


def test_cell_band_map_only_reads_cell_sheets():
    def payload(headers, rows):
        header_norm = {smart.norm_key(h): i for i, h in enumerate(headers)}
        return {
            "headers": headers,
            "header_index": {h: i for i, h in enumerate(headers)},
            "header_norm": header_norm,
            "rows": rows,
            "identity_idx": smart.identity_indexes(header_norm, fuzzy=False),
        }

    sheets = {
        "Cell": payload(
            ["*eNodeB Name", "*Local cell ID", "Frequency band"],
            [[None, 11, 8], [None, 14, 3], [None, 20, 1], [None, 74, 41]],
        ),
        "InterFreqHoGroup": payload(
            ["*eNodeB Name", "*Local cell ID", "AAAS Based Interfreq A1 RSRP Threshold(dBm)"],
            [[None, 11, -74], [None, 14, -118]],
        ),
        "NRDUCell": payload(
            ["*gNodeB Name", "*NR DU Cell ID", "*Frequency Band"],
            [[None, 102, "N41"]],
        ),
        "MAPPING DEF": payload(["SHEETNAME", "GROUP", "COLUMN", "MOC", "ATTR"], []),
    }

    class SpyCache:
        def __init__(self):
            self.got = []

        def names(self):
            return list(sheets)

        def get(self, name):
            self.got.append(name)
            return sheets.get(name)

    spy = SpyCache()
    index = audit.CellBandIndex(spy)
    assert "Cell" in spy.got
    assert "NRDUCell" in spy.got
    assert "InterFreqHoGroup" not in spy.got
    assert "MAPPING DEF" not in spy.got
    assert "11" in index.by_cell_id
    ho_row = sheets["InterFreqHoGroup"]
    ctx = smart.cell_from_row(ho_row["headers"], ho_row["header_norm"], ho_row["rows"][0], indexes=ho_row["identity_idx"])
    assert not ctx.band_tokens
    ctx = index.enrich(ctx)
    assert smart.band_matches("L09", ctx.band_tokens)

    class FakeCache:
        def get(self, name):
            return sheets[name]

    param = {
        "pid": "InterFreqHOA1ThdRsrp",
        "param_name": "AAAS Based Interfreq A1 RSRP Threshold",
        "recommend": "L9: -74 L18:-118 L21:-118 L26:-115",
        "mml": "InterFreqHoGroup",
    }
    resolved = {
        "reason": "OK",
        "sheet": "InterFreqHoGroup",
        "column": "AAAS Based Interfreq A1 RSRP Threshold",
        "bit": None,
        "smart_note": "",
    }
    out = audit.compare_param(param, resolved, FakeCache(), cell_index=index)
    assert out["status"] == "Consistent"
    assert out["match_count"] == 2
    assert out["mismatch_count"] == 0
    assert "L9" in (out.get("applied_recommend") or "") or "L09" in (out.get("applied_recommend") or "")
    hit, score, why = smart.closest_name("CelAlgoSwitch", ["CellAlgoSwitch", "CellMLB"], cutoff=0.8)
    assert hit == "CellAlgoSwitch"
    hit, score, why = smart.closest_name("NRDU CEL", ["NRDUCell", "Cell"], cutoff=0.8)
    assert smart.norm_key(hit) == "NRDUCELL"


def test_simple_compare_is_fast():
    """Simple recommend values must not scan Cell/band context on every row."""
    import time

    rows = [["MODIND", "ENODEBNAME", "LOCALCELLID", "SOMEPARAM"]]
    rows.append(["*eNodeB Name", "*eNodeB Name", "*Local cell ID", "Some Param"])
    for i in range(4000):
        rows.append([None, "SITE1", i, 1])
    payload_headers = ["*eNodeB Name", "*eNodeB Name", "*Local cell ID", "Some Param"]
    header_norm = {smart.norm_key(h): i for i, h in enumerate(payload_headers)}
    sheet = {
        "headers": payload_headers,
        "header_index": {h: i for i, h in enumerate(payload_headers)},
        "header_norm": header_norm,
        "rows": rows[2:],
        "identity_idx": smart.identity_indexes(header_norm, fuzzy=False),
    }

    class FakeCache:
        def get(self, _name):
            return sheet

    param = {
        "pid": "SomeParam",
        "param_name": "Some Param",
        "recommend": "1",
        "mml": "CellAlgoSwitch",
    }
    resolved = {
        "reason": "OK",
        "sheet": "CellAlgoSwitch",
        "column": "Some Param",
        "bit": None,
        "smart_note": "",
    }
    t0 = time.perf_counter()
    out = audit.compare_param(param, resolved, FakeCache(), cell_index=None)
    elapsed = time.perf_counter() - t0
    assert out["status"] == "Consistent"
    assert out["objects_checked"] == 4000
    assert elapsed < 1.5, f"simple compare too slow: {elapsed:.2f}s"


HO_ALLOWED_PACK = (
    "SnBasedInterFreqHoSw-0&GeranSepOpMobilitySwitch-0&UtranCsftbSwitch-0"
    "&GeranCsftbSwitch-1&UtranFlashCsftbSwitch-0&GeranFlashCsftbSwitch-0"
    "&CsftbAdaptiveBlindHoSwitch-1"
)


def _switch_sheet(rows):
    headers = ["*eNodeB Name", "*Local cell ID", "Handover Allowed Switch"]
    header_norm = {smart.norm_key(h): i for i, h in enumerate(headers)}
    return {
        "headers": headers,
        "header_index": {h: i for i, h in enumerate(headers)},
        "header_norm": header_norm,
        "rows": rows,
        "identity_idx": smart.identity_indexes(header_norm, fuzzy=False),
    }


def test_switch_bits_inside_parameter_column():
    """Recommend GeranCsftbSwitch-1 must be looked up inside the packed dump column."""
    assert audit.parse_switch_recommend("1") is None
    assert audit.parse_switch_recommend("ON") is None
    assert audit.parse_switch_recommend("-108") is None
    assert audit.parse_switch_recommend("L9:1") is None
    geran = audit.parse_switch_recommend("GeranCsftbSwitch-1")
    assert geran == [("GeranCsftbSwitch", "1")]
    multi = audit.parse_switch_recommend("GeranCsftbSwitch-1&UtranCsftbSwitch-0")
    assert multi == [("GeranCsftbSwitch", "1"), ("UtranCsftbSwitch", "0")]

    packed_ok = HO_ALLOWED_PACK
    packed_off = packed_ok.replace("GeranCsftbSwitch-1", "GeranCsftbSwitch-0")
    sheet = _switch_sheet(
        [
            [None, 11, packed_ok],
            [None, 14, packed_ok],
            [None, 20, packed_ok],
            [None, 74, packed_off],
        ]
    )

    class FakeCache:
        def get(self, _name):
            return sheet

    resolved = {
        "reason": "OK",
        "sheet": "CellAlgoSwitch",
        "column": "Handover Allowed Switch",
        "bit": None,
        "smart_note": "",
    }

    geran_param = {
        "pid": "HoAllowedSwitch",
        "param_name": "Handover Allowed Switch",
        "recommend": "GeranCsftbSwitch-1",
        "mml": "CELLALGOSWITCH",
    }
    out = audit.compare_param(geran_param, dict(resolved), FakeCache(), cell_index=None)
    assert out["status"] == "Mixed / Partial"
    assert out["match_count"] == 3
    assert out["mismatch_count"] == 1
    assert out["objects_checked"] == 4
    actuals = [name for name, _n in out["unique_actuals"]]
    assert "GeranCsftbSwitch-1" in actuals
    assert "GeranCsftbSwitch-0" in actuals
    assert not any("&SnBased" in a or "SnBasedInterFreqHoSw" in a for a in actuals)

    utran0 = dict(geran_param, recommend="UtranCsftbSwitch-0")
    out0 = audit.compare_param(utran0, dict(resolved), FakeCache(), cell_index=None)
    assert out0["status"] == "Consistent"
    assert out0["match_count"] == 4

    utran1 = dict(geran_param, recommend="UtranCsftbSwitch-1")
    out1 = audit.compare_param(utran1, dict(resolved), FakeCache(), cell_index=None)
    assert out1["status"] == "Inconsistent"
    assert out1["match_count"] == 0
    assert out1["mismatch_count"] == 4

    both = dict(geran_param, recommend="GeranCsftbSwitch-1&UtranCsftbSwitch-0")
    out_both = audit.compare_param(both, dict(resolved), FakeCache(), cell_index=None)
    assert out_both["match_count"] == 3
    assert out_both["mismatch_count"] == 1

    # Direct numeric values still compare the whole cell, not bits
    num_sheet = _switch_sheet([[None, 11, 20], [None, 14, 20]])
    num_sheet["headers"] = ["*eNodeB Name", "*Local cell ID", "Some Param"]
    num_sheet["header_index"] = {h: i for i, h in enumerate(num_sheet["headers"])}
    num_sheet["header_norm"] = {smart.norm_key(h): i for i, h in enumerate(num_sheet["headers"])}

    class NumCache:
        def get(self, _name):
            return num_sheet

    num_param = {
        "pid": "SomeParam",
        "param_name": "Some Param",
        "recommend": "20",
        "mml": "CellMLB",
    }
    num_resolved = {
        "reason": "OK",
        "sheet": "CellMLB",
        "column": "Some Param",
        "bit": None,
        "smart_note": "",
    }
    num_out = audit.compare_param(num_param, num_resolved, NumCache(), cell_index=None)
    assert num_out["status"] == "Consistent"
    assert num_out["match_count"] == 2


def test_conditions_group_id_filter():
    """Conditions Interfreq handover group ID=0 audits only matching dump rows."""
    wrapped = "Interfreq handover group\nID=0"
    clauses = smart.parse_conditions(wrapped)
    assert len(clauses) == 1
    assert clauses[0].value == "0"
    assert smart.norm_key(clauses[0].display_name) == "INTERFREQHANDOVERGROUPID"

    named = smart.parse_conditions("Interfreq handover group ID (INTERFREQHOGROUPID)=1")
    assert named[0].value == "1"
    assert named[0].short_name == "INTERFREQHOGROUPID"
    assert "INTERFREQHOGROUPID" in smart.condition_name_keys(named[0])
    assert smart.canon_id_value(0.0) == "0"
    assert smart.canon_id_value(1) == "1"

    headers = [
        "*eNodeB Name",
        "*Local cell ID",
        "*Interfreq handover group ID",
        "AAAS Based Interfreq A1 RSRP Threshold(dBm)",
        "Load Based Interfreq RSRP threshold",
    ]
    # One Local cell ID exists twice: group 0 = Data, group 1 = VoLTE
    rows = [
        ["TNMdp05", 11, 0, -74, -103],
        ["TNMdp05", 11, 1, -90, -105],
        ["TNMdp05", 14, 0.0, -118, -103],
        ["TNMdp05", 14, 1, -90, -105],
        ["TNMdp05", 20, 0, -118, -103],
        ["TNMdp05", 20, 1, -90, -105],
        ["TNMdp05", 74, 0, -115, -103],
        ["TNMdp05", 74, 1, -90, -105],
    ]
    header_norm = {smart.norm_key(h): i for i, h in enumerate(headers)}
    sheet = {
        "headers": headers,
        "header_index": {h: i for i, h in enumerate(headers)},
        "header_norm": header_norm,
        "rows": rows,
        "identity_idx": smart.identity_indexes(header_norm, fuzzy=False),
    }

    class FakeCache:
        def get(self, _name):
            return sheet

    cell_payload = {
        "headers": ["*eNodeB Name", "*Local cell ID", "Frequency band"],
        "header_index": {},
        "header_norm": {
            "ENODEBNAME": 0,
            "LOCALCELLID": 1,
            "FREQUENCYBAND": 2,
        },
        "rows": [
            ["TNMdp05", 11, 8],
            ["TNMdp05", 14, 3],
            ["TNMdp05", 20, 1],
            ["TNMdp05", 74, 41],
        ],
        "identity_idx": smart.identity_indexes(
            {"ENODEBNAME": 0, "LOCALCELLID": 1, "FREQUENCYBAND": 2}, fuzzy=False
        ),
    }
    cell_payload["header_index"] = {h: i for i, h in enumerate(cell_payload["headers"])}

    class BandCache:
        def names(self):
            return ["Cell", "InterFreqHoGroup"]

        def get(self, name):
            if name == "Cell":
                return cell_payload
            return sheet

    index = audit.CellBandIndex(BandCache())
    resolved = {
        "reason": "OK",
        "sheet": "InterFreqHoGroup",
        "column": "AAAS Based Interfreq A1 RSRP Threshold",
        "bit": None,
        "smart_note": "",
    }
    a1 = {
        "pid": "InterFreqHOA1ThdRsrp",
        "param_name": "AAAS Based Interfreq A1 RSRP Threshold",
        "recommend": "L9: -74 L18:-118 L21:-118 L26:-115",
        "conditions": "Interfreq handover group ID=0",
        "mml": "InterFreqHoGroup",
    }
    out0 = audit.compare_param(a1, dict(resolved), FakeCache(), cell_index=index)
    assert out0["status"] == "Consistent"
    assert out0["objects_checked"] == 4
    assert out0["match_count"] == 4
    assert out0["mismatch_count"] == 0
    assert out0["skipped_not_applicable"] == 4

    a1_volte = dict(a1, conditions="Interfreq handover group ID (INTERFREQHOGROUPID)=1")
    out1 = audit.compare_param(a1_volte, dict(resolved), FakeCache(), cell_index=index)
    assert out1["objects_checked"] == 4
    assert out1["match_count"] == 0
    assert out1["mismatch_count"] == 4
    assert out1["status"] == "Inconsistent"

    none = dict(a1, conditions=None)
    out_all = audit.compare_param(none, dict(resolved), FakeCache(), cell_index=index)
    assert out_all["objects_checked"] == 8

    load_resolved = dict(resolved, column="Load Based Interfreq RSRP threshold")
    load = {
        "pid": "InterFreqLoadBasedHoA4ThdRsrp",
        "param_name": "Load Based Interfreq RSRP threshold",
        "recommend": "-105",
        "conditions": "Interfreq handover group ID (INTERFREQHOGROUPID)=1",
        "mml": "InterFreqHoGroup",
    }
    load_out = audit.compare_param(load, load_resolved, FakeCache(), cell_index=None)
    assert load_out["objects_checked"] == 4
    assert load_out["match_count"] == 4
    assert load_out["status"] == "Consistent"

    missing = dict(a1, conditions="Interfreq handover group ID=9")
    miss_out = audit.compare_param(missing, dict(resolved), FakeCache(), cell_index=index)
    assert miss_out["objects_checked"] == 0
    assert miss_out["status"] == "No Recommend Value"
    assert "Conditions" in (miss_out["remark"] or "")


def test_list_workbooks_rats():
    with tempfile.TemporaryDirectory() as raw:
        root = Path(raw)
        audit.ensure_rat_input_folders(root)
        (root / "4G" / "lte.xlsx").write_bytes(b"")  # not a real xlsx, is_workbook checks suffix
        # is_workbook only checks suffix; empty file still counts as workbook by extension
        (root / "5G" / "nr.xlsx").write_bytes(b"")
        (root / "loose.xlsx").write_bytes(b"")
        only_4g = audit.list_workbooks(root, rats=["4G"], use_rat_subfolders=True)
        assert [p.name for p in only_4g] == ["lte.xlsx"]
        only_5g = audit.list_workbooks(root, rats=["5G"], use_rat_subfolders=True)
        assert [p.name for p in only_5g] == ["nr.xlsx"]
        all_rats = audit.list_workbooks(root, rats=["5G", "4G", "3G", "2G"], use_rat_subfolders=True)
        names = {p.name for p in all_rats}
        assert names == {"lte.xlsx", "nr.xlsx", "loose.xlsx"}


def _build_4g_dump(path: Path):
    wb = Workbook()
    _write_sheet(
        wb,
        "Cell",
        [
            ["MODIND", "ENODEBNAME", "LOCALCELLID", "FREQUENCYBAND"],
            ["*eNodeB Name", "*eNodeB Name", "*Local cell ID", "Frequency band"],
            [None, "DHAPT08", 11, 8],
            [None, "DHAPT08", 14, 3],
            [None, "DHAPT08", 20, 1],
            [None, "DHAPT08", 74, 41],
        ],
    )
    _write_sheet(
        wb,
        "CellAlgoSwitch",
        [
            ["MODIND", "ENODEBNAME", "LOCALCELLID", "ENHANCEDMLBALGOSWITCH", "HOALLOWEDSWITCH"],
            [
                "*eNodeB Name",
                "*eNodeB Name",
                "*Local cell ID",
                "Enhanced MLB Algorithm Switch",
                "Handover Allowed Switch",
            ],
            [
                None,
                "DHAPT08",
                11,
                "SpectralETBasedLoadEvalSw-1&OtherSw-0",
                "SnBasedInterFreqHoSw-0&GeranSepOpMobilitySwitch-0&UtranCsftbSwitch-0&GeranCsftbSwitch-1&UtranFlashCsftbSwitch-0&GeranFlashCsftbSwitch-0&CsftbAdaptiveBlindHoSwitch-1",
            ],
            [
                None,
                "DHAPT08",
                14,
                "SpectralETBasedLoadEvalSw-1&OtherSw-0",
                "SnBasedInterFreqHoSw-0&GeranSepOpMobilitySwitch-0&UtranCsftbSwitch-0&GeranCsftbSwitch-1&UtranFlashCsftbSwitch-0&GeranFlashCsftbSwitch-0&CsftbAdaptiveBlindHoSwitch-1",
            ],
            [
                None,
                "DHAPT08",
                20,
                "SpectralETBasedLoadEvalSw-0&OtherSw-0",
                "SnBasedInterFreqHoSw-0&GeranSepOpMobilitySwitch-0&UtranCsftbSwitch-0&GeranCsftbSwitch-1&UtranFlashCsftbSwitch-0&GeranFlashCsftbSwitch-0&CsftbAdaptiveBlindHoSwitch-1",
            ],
            [
                None,
                "DHAPT08",
                74,
                "SpectralETBasedLoadEvalSw-1&OtherSw-0",
                "SnBasedInterFreqHoSw-0&GeranSepOpMobilitySwitch-0&UtranCsftbSwitch-0&GeranCsftbSwitch-0&UtranFlashCsftbSwitch-0&GeranFlashCsftbSwitch-0&CsftbAdaptiveBlindHoSwitch-1",
            ],
        ],
    )
    _write_sheet(
        wb,
        "CellMLB",
        [
            ["MODIND", "ENODEBNAME", "LOCALCELLID", "INTERFREQIDLEMLBUENUMTHD"],
            ["*eNodeB Name", "*eNodeB Name", "*Local cell ID", "Inter-Freq Idle MLB UE Number Threshold"],
            [None, "DHAPT08", 11, 1],
            [None, "DHAPT08", 14, 20],
            [None, "DHAPT08", 20, 20],
            [None, "DHAPT08", 74, 99],
        ],
    )
    _write_sheet(
        wb,
        "InterFreqHoGroup",
        [
            ["MODIND", "ENODEBNAME", "LOCALCELLID", "INTERFREQHOGROUPID", "A3OFFSET", "INTERFREQHOA1THDRSRP", "INTERFREQHOA2THDRSRP"],
            [
                "*eNodeB Name",
                "*eNodeB Name",
                "*Local cell ID",
                "*Interfreq handover group ID",
                "A3 Offset",
                "AAAS Based Interfreq A1 RSRP Threshold(dBm)",
                "AAAS Based Interfreq A2 RSRP Threshold(dBm)",
            ],
            [None, "DHAPT08", 11, 0, -108, -74, -99],
            [None, "DHAPT08", 11, 1, -108, -74, -99],
            [None, "DHAPT08", 14, 0, -108, -118, -99],
            [None, "DHAPT08", 14, 1, -108, -118, -99],
            [None, "DHAPT08", 20, 0, -108, -118, -99],
            [None, "DHAPT08", 20, 1, -100, -118, -99],
            [None, "DHAPT08", 74, 0, -108, -115, -99],
            [None, "DHAPT08", 74, 1, -108, -115, -99],
        ],
    )
    _save(wb, path)


def _build_5g_dump(path: Path):
    wb = Workbook()
    _write_sheet(
        wb,
        "NRDUCell",
        [
            ["MODIND", "GNODEBNAME", "NRDUCELLID", "NRDUCELLNAME", "FREQUENCYBAND"],
            ["*gNodeB Name", "*gNodeB Name", "*NR DU Cell ID", "*NR DU Cell Name", "*Frequency Band"],
            [None, "TNMDP05", 102, "TNMDP05N26B", "N41"],
            [None, "GPSDRP7", 101, "GPSDRP7N26A", "N41"],
        ],
    )
    _write_sheet(
        wb,
        "NRDUCellAlgoSwitch",
        [
            ["MODIND", "GNODEBNAME", "NRDUCELLID", "SOMEALGOSWITCH"],
            ["*gNodeB Name", "*gNodeB Name", "*NR DU Cell ID", "Some Algo Switch"],
            [None, "TNMDP05", 102, "DemoSw-1"],
            [None, "GPSDRP7", 101, "DemoSw-1"],
        ],
    )
    _save(wb, path)


def _build_reference(path: Path):
    wb = Workbook()
    _write_sheet(
        wb,
        "MLB Plan",
        [
            ["MO Name", "Parameter ID", "Parameter Name", "Default Value", "Proposed Value"],
            [
                "CellAlgoSwitch",
                "SpectralETBasedLoadEvalSw@EnhancedMlbAlgoSwitch",
                "Enhanced MLB Algorithm Switch",
                "Off",
                "L9: 1\nL18: 1\nL21: 1\nL26: 1",
            ],
            [
                "Cell MLB",  # spaced / informal MO name — smart matching should recover CellMLB
                "InterFreqIdleMlbUeNumThd",
                "Inter-Freq Idle MLB UE Number Threshold",
                "100",
                "L9: 1\nL18: 20\nL21: 20\nL26: 35",
            ],
            [
                "InterFreqHoGroup",
                "A3Offset",
                "A3 Offset",
                "",
                "(InterFreqHoGroupId=1)=>L09=-108\n(InterFreqHoGroupId=1)=>L18=-108\n(InterFreqHoGroupId=1)=>L21=-108",
            ],
            [
                "InterFreqHoGroup",
                "InterFreqHOA1ThdRsrp",
                "AAAS Based Interfreq A1 RSRP Threshold",
                "",
                "L9: -74 L18:-118 L21:-118 L26:-115",
            ],
        ],
    )
    _write_sheet(
        wb,
        "Mobility",
        [
            [
                "MO Name",
                "Parameter ID",
                "Parameter Name",
                "Default Value",
                "Proposed Value",
                "Conditions",
                "Targ. t/Impact",
            ],
            [
                "InterFreqHoGroup",
                "InterFreqHOA1ThdRsrp",
                "AAAS Based Interfreq A1 RSRP Threshold",
                "-105",
                "L9: -74 L18:-118 L21:-118 L26:-115",
                "Interfreq handover group ID=0",
                "Mobility",
            ],
            [
                "InterFreqHoGroup",
                "A3Offset",
                "A3 Offset",
                "",
                "-108",
                "Interfreq handover group ID (INTERFREQHOGROUPID)=1",
                "Mobility",
            ],
        ],
    )
    _write_sheet(
        wb,
        "Missing Name",
        [
            ["MO Name", "Parameter ID", "Default Value", "Proposed Value"],
            ["CellMLB", "Something@OtherSwitch", "Off", "1"],
        ],
    )
    _write_sheet(
        wb,
        "Legacy",
        [
            ["Function", "MML Object", "Parameter ID", "Recommend Value"],
            ["Idle", "CellMLB", "InterFreqIdleMlbUeNumThd", "20"],
        ],
    )
    _write_sheet(
        wb,
        "Three Keys",
        [
            ["MML Object", "Parameter ID", "Parameter Name", "Golden"],
            [
                "CellAlgoSwitch",
                "SpectralETBasedLoadEvalSw@EnhancedMlbAlgoSwitch",
                "Enhanced MLB Algorithm Switch",
                "1",
            ],
            [
                "CELLALGOSWITCH",
                "HoAllowedSwitch",
                "Handover Allowed Switch",
                "GeranCsftbSwitch-1",
            ],
            [
                "CELLALGOSWITCH",
                "HoAllowedSwitch",
                "Handover Allowed Switch",
                "UtranCsftbSwitch-0",
            ],
            [
                "CELLALGOSWITCH",
                "HoAllowedSwitch",
                "Handover Allowed Switch",
                "UtranCsftbSwitch-1",
            ],
        ],
    )
    _save(wb, path)


def test_end_to_end_audit():
    key = lic.find_private_key_path()
    assert key is not None, "private key required for licensed audit test"
    with tempfile.TemporaryDirectory() as raw:
        root = Path(raw)
        input_dir = root / "Input"
        ref_dir = root / "Reference"
        out_dir = root / "Output"
        _build_4g_dump(input_dir / "4G" / "lte_dump.xlsx")
        _build_5g_dump(input_dir / "5G" / "nr_dump.xlsx")
        _build_reference(ref_dir / "plan.xlsx")
        license_path = root / "test.lic"
        lic.write_license(lic.issue_license("Test", days=7, private_key_path=key), license_path)

        only_4g = audit.list_workbooks(input_dir, rats=["4G"], use_rat_subfolders=True)
        assert [p.name for p in only_4g] == ["lte_dump.xlsx"]
        only_5g = audit.list_workbooks(input_dir, rats=["5G"], use_rat_subfolders=True)
        assert [p.name for p in only_5g] == ["nr_dump.xlsx"]

        run, summary, _extras = audit.execute_folder_audit(
            output_dir=out_dir,
            input_folder=input_dir,
            reference_folder=ref_dir,
            license_file=license_path,
            rats=["4G"],
            progress=lambda *_a, **_k: None,
        )
        params = _params_from_csv(run.out_csv)
        mlb_switch = next(p for p in params if "SpectralETBasedLoadEvalSw" in p["parameter_id"])
        assert mlb_switch["status"] in {"Mixed / Partial", "Inconsistent"}
        # L21 cell (local 20) has bit 0, others 1 → mixed
        assert int(mlb_switch["mismatch_count"]) >= 1
        assert int(mlb_switch["match_count"]) >= 1

        thd = next(p for p in params if p["parameter_id"] == "InterFreqIdleMlbUeNumThd" and p["reference_sheet"] == "MLB Plan")
        assert thd["status"] in {"Mixed / Partial", "Inconsistent"}
        # L26 cell 74 has 99 vs expected 35
        assert int(thd["mismatch_count"]) >= 1

        missing = next(p for p in params if p["reference_sheet"] == "Missing Name")
        assert missing["status"] == "No Recommend Value"
        assert "Parameter Name" in missing["remark"]

        group = next(p for p in params if p["parameter_id"] == "A3Offset")
        assert int(group["objects_checked"]) >= 1
        # group 0 cell skipped; L21 group 1 has -100 vs -108
        assert int(group["mismatch_count"]) >= 1

        a1 = next(p for p in params if p["parameter_id"] == "InterFreqHOA1ThdRsrp" and p["reference_sheet"] == "MLB Plan")
        assert a1["status"] == "Consistent"
        assert int(a1["match_count"]) == 8
        assert int(a1["mismatch_count"]) == 0
        assert "A2" not in (a1.get("config_column") or "")

        cond0 = next(
            p
            for p in params
            if p["parameter_id"] == "InterFreqHOA1ThdRsrp" and p["reference_sheet"] == "Mobility"
        )
        assert int(cond0["objects_checked"]) == 4
        assert cond0["status"] == "Consistent"
        assert "Interfreq handover group ID=0" in (cond0.get("conditions") or "")

        cond_a3 = next(
            p
            for p in params
            if p["parameter_id"] == "A3Offset" and p["reference_sheet"] == "Mobility"
        )
        assert int(cond_a3["objects_checked"]) == 4
        # group-1 L21 cell 20 is -100 vs -108; other group-1 cells match
        assert int(cond_a3["mismatch_count"]) >= 1
        assert int(cond_a3["match_count"]) >= 1

        legacy = next(p for p in params if p["reference_sheet"] == "Legacy")
        assert legacy["status"] in {"Mixed / Partial", "Inconsistent", "Consistent"}

        three = next(p for p in params if p["reference_sheet"] == "Three Keys" and "SpectralETBasedLoadEvalSw" in p["parameter_id"])
        assert three["parameter_name"] == "Enhanced MLB Algorithm Switch"
        assert three["status"] in {"Mixed / Partial", "Inconsistent", "Consistent"}
        assert int(three["objects_checked"]) >= 1

        geran = next(
            p
            for p in params
            if p["reference_sheet"] == "Three Keys" and p["recommend_value"] == "GeranCsftbSwitch-1"
        )
        assert geran["status"] == "Mixed / Partial"
        assert int(geran["match_count"]) == 3
        assert int(geran["mismatch_count"]) == 1
        assert "GeranCsftbSwitch" in (geran.get("bit_name") or "")

        utran_off = next(
            p
            for p in params
            if p["reference_sheet"] == "Three Keys" and p["recommend_value"] == "UtranCsftbSwitch-0"
        )
        assert utran_off["status"] == "Consistent"
        assert int(utran_off["match_count"]) == 4

        utran_on = next(
            p
            for p in params
            if p["reference_sheet"] == "Three Keys" and p["recommend_value"] == "UtranCsftbSwitch-1"
        )
        assert utran_on["status"] == "Inconsistent"
        assert int(utran_on["match_count"]) == 0
        assert int(utran_on["mismatch_count"]) == 4

        assert run.out_xlsx.exists()
        assert "4G" in ",".join(run.selected_rats)


def _params_from_csv(path: Path):
    import csv

    with path.open(encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def main() -> int:
    test_rats_and_split()
    test_band_families()
    test_recommend_parser()
    test_three_identity_columns()
    test_fuzzy_names()
    test_a1_not_matched_as_a2()
    test_cell_band_map_only_reads_cell_sheets()
    test_simple_compare_is_fast()
    test_switch_bits_inside_parameter_column()
    test_conditions_group_id_filter()
    test_list_workbooks_rats()
    test_end_to_end_audit()
    print("rat-band-smart-audit tests passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
