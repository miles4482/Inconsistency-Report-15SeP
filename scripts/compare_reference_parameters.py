#!/usr/bin/env python3
"""Compare every reference workbook against every input dump.

Three local folders:
    Input Folder      configuration dumps (any names, any RAT, many files)
    Reference Folder  recommended / plan values (any names, many files)
    Output Folder     inconsistency reports

Regular use:
    python3 scripts/parameter_audit_app.py
    python3 scripts/compare_reference_parameters.py --input-folder ... --reference-folder ... --output-folder ... --rats 4G,5G
    ./run_audit.sh
    run_audit.bat
"""

from __future__ import annotations

import argparse
import csv
import fnmatch
import json
import re
import shutil
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pyxlsb import open_workbook

import license_control as license_mod
import smart_match as smart
from app_version import APP_TITLE, APP_VERSION


WORKBOOK_SUFFIXES = {".xlsx", ".xlsb", ".xlsm"}


def runtime_root() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent.parent


ROOT = runtime_root()
DEFAULT_INPUT_DIR = ROOT / "input"
DEFAULT_REFERENCE_DIR = ROOT / "reference"
DEFAULT_OUTPUT_DIR = ROOT / "reports"
DEFAULT_CONFIG = ROOT / "config" / "audit.json"

REF_PATTERNS = (
    "Reference Parameter*.xlsx",
    "Reference Parameter*.xlsb",
    "*Reference*Parameter*.xlsx",
)
CFG_4G_PATTERNS = (
    "4G_ConfigurationData*.xlsb",
    "4G_ConfigurationData*.xlsx",
    "*4G*Config*.xlsb",
    "*4G*Config*.xlsx",
)
CFG_5G_PATTERNS = (
    "5G_ConfigurationData*.xlsb",
    "5G_ConfigurationData*.xlsx",
    "*5G*Config*.xlsb",
    "*5G*Config*.xlsx",
)

TRUE_SET = {"1", "1.0", "ON", "TRUE", "YES", "ENABLE", "ENABLED"}
FALSE_SET = {"0", "0.0", "OFF", "FALSE", "NO", "DISABLE", "DISABLED"}
UNIT_RE = re.compile(r"^(-?\d+(?:\.\d+)?)(MS|S|DBM|DB|MHZ|KHZ|MIN|DAY)?$", re.I)
SWITCH_ASSIGN_RE = re.compile(
    r"^([A-Za-z][A-Za-z0-9_]*)\s*-\s*(0|1|ON|OFF|TRUE|FALSE|YES|NO)$",
    re.I,
)
META_SHEETS = {
    "SUMMARYRES",
    "FILEIDENTIFICATION",
    "INNERVALIDEDEF",
    "MAPPINGCELLTEMPLATE",
    "PRODUCTTYPE",
    "INNER_VERSION_FOR_VALID",
    "CONTROL INFOR",
    "CONTROL DEF",
    "SHEET DEF",
    "MAPPING DEF",
    "VALID DEF",
    "COMMENTS",
}
SKIP_REF_SHEETS = {
    "TITLE",
    "SUMMARY",
    "COVER",
    "INDEX",
    "README",
    "REVISION",
    "CHANGELOG",
    "HISTORY",
    "NOTES",
}
REF_HEADER_ALIASES = {
    "function": {"FUNCTION", "FEATURE", "CATEGORY", "AREA", "DOMAIN", "GROUP"},
    "mml": {
        "MMLOBJECT",
        "MMLOBJECTNAME",
        "MML",
        "MO",
        "MOC",
        "MONAME",
        "OBJECT",
        "OBJECTNAME",
        "MANAGEDOBJECT",
        "MMLOBJECTNAME",
        "NEOBJECT",
    },
    "pid": {
        "PARAMETERID",
        "PARAMETER",
        "PARAMID",
        "PARA",
        "ATTR",
        "ATTRIBUTE",
        "PARAM",
    },
    "param_name": {
        "PARAMETERNAME",
        "PARAMNAME",
        "PARANAME",
        "DISPLAYNAME",
        "FULLNAME",
        "ATTRIBUTENAME",
    },
    "recommend": {
        "RECOMMENDVALUE",
        "RECOMMENDEDVALUE",
        "RECOMMEND",
        "RECOMMENDED",
        "PROPOSEDVALUE",
        "PROPOSED",
        "PLANVALUE",
        "PLANNEDVALUE",
        "TARGETVALUE",
        "EXPECTEDVALUE",
        "BASELINEVALUE",
        "REFVALUE",
        "STANDARDVALUE",
        "GOLDENVALUE",
        "PLAN",
    },
}

# Microsoft Excel worksheet limits. This tool does not impose a lower cap.
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLS = 16_384


@dataclass
class RunContext:
    root: Path
    reference: Path
    cfg_4g: Path | None
    cfg_5g: Path | None
    output_dir: Path
    run_id: str
    audit_date: str
    out_xlsx: Path
    out_md: Path
    out_csv: Path
    out_summary: Path
    history_csv: Path
    latest_dir: Path
    input_folder: Path | None = None
    reference_folder: Path | None = None
    input_files: list = field(default_factory=list)
    reference_files: list = field(default_factory=list)
    selected_rats: list = field(default_factory=lambda: list(smart.RAT_FOLDERS))


def workbook_kind(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".xlsb":
        return "xlsb"
    if suffix in {".xlsx", ".xlsm"}:
        return "xlsx"
    raise ValueError(f"Unsupported workbook type: {path.name} ({suffix or 'no extension'}). Use .xlsx, .xlsb or .xlsm.")


def is_workbook(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() in WORKBOOK_SUFFIXES


def ensure_rat_input_folders(folder: Path | None):
    if folder is None:
        return
    folder = Path(folder).expanduser()
    folder.mkdir(parents=True, exist_ok=True)
    for rat in smart.RAT_FOLDERS:
        (folder / rat).mkdir(parents=True, exist_ok=True)


def list_workbooks(folder: Path | None, rats=None, use_rat_subfolders: bool = False) -> list[Path]:
    if folder is None:
        return []
    folder = Path(folder).expanduser()
    if not folder.exists() or not folder.is_dir():
        return []
    if not use_rat_subfolders:
        files = [p for p in folder.iterdir() if is_workbook(p)]
        files.sort(key=lambda p: p.name.lower())
        return files
    selected = smart.normalize_rats(rats)
    files = []
    seen = set()
    for rat in selected:
        for name in smart.rat_subfolder_names(rat):
            sub = folder / name
            if not sub.is_dir():
                continue
            for path in sub.iterdir():
                key = path.resolve() if path.exists() else path
                if is_workbook(path) and key not in seen:
                    files.append(path)
                    seen.add(key)
    if set(selected) == set(smart.RAT_FOLDERS):
        for path in folder.iterdir():
            key = path.resolve() if path.exists() else path
            if is_workbook(path) and key not in seen:
                files.append(path)
                seen.add(key)
    files.sort(key=lambda p: (p.parent.name.lower(), p.name.lower()))
    return files


def workbook_rel_label(path: Path, folder: Path | None) -> str:
    path = Path(path)
    if folder:
        try:
            rel = path.relative_to(Path(folder).expanduser().resolve())
            return str(rel).replace("\\", "/")
        except ValueError:
            pass
    if path.parent.name in smart.RAT_FOLDERS:
        return f"{path.parent.name}/{path.name}"
    return path.name


def sheet_names(path: Path) -> list[str]:
    if workbook_kind(path) == "xlsb":
        with open_workbook(str(path)) as wb:
            return list(wb.sheets)
    wb = load_workbook(path, read_only=True, data_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def read_sheet_rows(path: Path, sheet_name: str):
    if workbook_kind(path) == "xlsb":
        with open_workbook(str(path)) as wb:
            if sheet_name not in wb.sheets:
                return None
            with wb.get_sheet(sheet_name) as sheet:
                return [[c.v for c in row] for row in sheet.rows()]
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return None
    rows = [list(row) for row in wb[sheet_name].iter_rows(values_only=True)]
    wb.close()
    return rows


def unique_files(paths):
    seen = {}
    for path in paths:
        if path.is_file():
            seen[path.resolve()] = path
    return list(seen.values())


def find_files(search_dirs, patterns) -> list[Path]:
    found = []
    for folder in search_dirs:
        if not folder or not folder.exists():
            continue
        for pattern in patterns:
            found.extend(folder.glob(pattern))
    return unique_files(found)


def newest_file(paths: list[Path]) -> Path | None:
    if not paths:
        return None
    return max(paths, key=lambda p: (p.stat().st_mtime, p.name))


def load_audit_config(path: Path | None) -> dict:
    if path is None or not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def search_dirs(root: Path, input_dir: Path) -> list[Path]:
    dirs = [input_dir, root]
    unique = []
    seen = set()
    for folder in dirs:
        key = folder.resolve() if folder.exists() else folder
        if key not in seen:
            unique.append(folder)
            seen.add(key)
    return unique


def discover_inputs(root: Path, input_dir: Path, config: dict | None = None) -> dict:
    dirs = search_dirs(root, input_dir)
    patterns = (config or {}).get("file_patterns") or {}
    ref_patterns = tuple(patterns.get("reference") or REF_PATTERNS)
    cfg_4g_patterns = tuple(patterns.get("4G") or CFG_4G_PATTERNS)
    cfg_5g_patterns = tuple(patterns.get("5G") or CFG_5G_PATTERNS)
    return {
        "reference": newest_file(find_files(dirs, ref_patterns)),
        "cfg_4g": newest_file(find_files(dirs, cfg_4g_patterns)),
        "cfg_5g": newest_file(find_files(dirs, cfg_5g_patterns)),
        "search_dirs": dirs,
    }


def _matches_any(name: str, patterns) -> bool:
    return any(fnmatch.fnmatch(name, pat) for pat in patterns)


def guess_file_role(path: Path) -> str:
    name = path.name
    upper = name.upper()
    if _matches_any(name, REF_PATTERNS) or "REFERENCE" in upper:
        return "reference"
    if _matches_any(name, CFG_4G_PATTERNS) or re.search(r"(^|[^0-9])4G([^0-9]|$)", upper):
        return "4g"
    if _matches_any(name, CFG_5G_PATTERNS) or re.search(r"(^|[^0-9])5G([^0-9]|$)", upper):
        return "5g"
    try:
        sheets = [str(s).upper() for s in sheet_names(path)]
    except Exception:
        return "unknown"
    if any(s in {"NR PERFORMANCE", "NR ANCHOR", "TITLE"} for s in sheets) and "MAPPING DEF" not in sheets:
        return "reference"
    if "NR PERFORMANCE" in sheets or "NR ANCHOR" in sheets:
        return "reference"
    blob = " ".join(sheets)
    if "NRCELL" in blob or "GNODEBPARAM" in blob or "GNBX2SONCONFIG" in blob:
        return "5g"
    if "NSADCMGMTCONFIG" in blob or "ENODEBALGOSWITCH" in blob or "CELLALGOSWITCH" in blob:
        return "4g"
    return "unknown"


def classify_input_files(paths) -> dict:
    classified = {"reference": [], "4g": [], "5g": [], "unknown": []}
    for raw in paths:
        path = Path(raw).expanduser().resolve()
        if not path.is_file():
            classified["unknown"].append(path)
            continue
        classified[guess_file_role(path)].append(path)
    picked = {
        "reference": newest_file(classified["reference"]),
        "cfg_4g": newest_file(classified["4g"]),
        "cfg_5g": newest_file(classified["5g"]),
        "unknown": classified["unknown"],
        "all_reference": classified["reference"],
        "all_4g": classified["4g"],
        "all_5g": classified["5g"],
    }
    return picked


def build_run_context(
    root: Path,
    reference: Path,
    cfg_4g: Path | None,
    cfg_5g: Path | None,
    output_dir: Path,
    run_id: str | None = None,
    input_folder: Path | None = None,
    reference_folder: Path | None = None,
    input_files: list | None = None,
    reference_files: list | None = None,
    selected_rats: list | None = None,
) -> RunContext:
    now = datetime.now()
    run_id = run_id or now.strftime("%Y%m%d_%H%M%S")
    run_dir = output_dir / "runs" / run_id
    run_dir.mkdir(parents=True, exist_ok=True)
    latest_dir = output_dir / "latest"
    latest_dir.mkdir(parents=True, exist_ok=True)
    reference_files = [Path(p) for p in (reference_files or [reference])]
    input_files = [Path(p) for p in (input_files or [p for p in (cfg_4g, cfg_5g) if p])]
    return RunContext(
        root=root,
        reference=Path(reference),
        cfg_4g=Path(cfg_4g) if cfg_4g else None,
        cfg_5g=Path(cfg_5g) if cfg_5g else None,
        output_dir=output_dir,
        run_id=run_id,
        audit_date=now.strftime("%d %b %Y"),
        out_xlsx=run_dir / "Parameter_Inconsistency_Report.xlsx",
        out_md=run_dir / "Parameter_Inconsistency_Report.md",
        out_csv=run_dir / "all_parameters.csv",
        out_summary=run_dir / "run_summary.json",
        history_csv=output_dir / "run_history.csv",
        latest_dir=latest_dir,
        input_folder=Path(input_folder) if input_folder else None,
        reference_folder=Path(reference_folder) if reference_folder else None,
        input_files=input_files,
        reference_files=reference_files,
        selected_rats=list(selected_rats or smart.RAT_FOLDERS),
    )


def copy_latest(run: RunContext):
    for src, name in (
        (run.out_xlsx, "Parameter_Inconsistency_Report.xlsx"),
        (run.out_md, "Parameter_Inconsistency_Report.md"),
        (run.out_csv, "all_parameters.csv"),
        (run.out_summary, "run_summary.json"),
    ):
        if src.exists():
            shutil.copy2(src, run.latest_dir / name)


def norm_key(value) -> str:
    if value is None:
        return ""
    return re.sub(r"[^A-Z0-9]", "", str(value).upper())


def clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def is_empty(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


def canon_value(value):
    if is_empty(value):
        return None
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, float):
        if value.is_integer():
            value = int(value)
        else:
            return str(value)
    if isinstance(value, int):
        return str(value)
    text = clean_text(value)
    compact = text.upper().replace(" ", "")
    unit_match = UNIT_RE.fullmatch(compact)
    if unit_match:
        number = float(unit_match.group(1))
        if number.is_integer():
            return str(int(number))
        return str(number)
    if compact in TRUE_SET:
        return "1"
    if compact in FALSE_SET:
        return "0"
    return re.sub(r"[^A-Z0-9]", "", compact)


def values_match(actual, recommend) -> bool:
    if is_empty(recommend):
        return False
    left = canon_value(actual)
    right = canon_value(recommend)
    if left is None or right is None:
        return False
    return left == right


def parse_bitfield(value):
    if is_empty(value):
        return None
    text = clean_text(value)
    if "&" not in text and not re.search(r"-\s*[01]\b", text):
        return None
    bits = {}
    for part in text.split("&"):
        part = part.strip()
        if not part:
            continue
        name, sep, raw = part.rpartition("-")
        if not sep:
            continue
        bits[norm_key(name)] = raw.strip()
    return bits or None


def extract_bit(value, bit_name):
    bits = parse_bitfield(value)
    if bits is None:
        return None, "not_bitfield"
    key = norm_key(bit_name)
    if key in bits:
        return bits[key], "found"
    return None, "bit_missing"


def parse_switch_recommend(value):
    """Recommend that names switch/bit options inside a packed parameter.

    Direct values (1, ON, -108, L9:1) return None.
    Switch recommends (GeranCsftbSwitch-1 or A-0&B-1) return
    a list of (bit_name, expected_value).
    """
    text = clean_text(value)
    if not text:
        return None
    if smart.BAND_VALUE_RE.search(text) or (text.startswith("(") and "=" in text):
        return None
    parts = [p.strip() for p in text.split("&") if p.strip()]
    if not parts:
        return None
    parsed = []
    for part in parts:
        match = SWITCH_ASSIGN_RE.fullmatch(part)
        if not match:
            return None
        parsed.append((match.group(1), match.group(2)))
    if len(parsed) == 1 and len(parsed[0][0]) < 3:
        return None
    return parsed


def audit_switch_bits(raw, expected_pairs):
    """Look up each recommended switch/bit inside the dump parameter column.

    Extra switches present in the dump are ignored. Returns
    (display, state, all_match, missing_names).
    """
    dump_bits = parse_bitfield(raw)
    if not dump_bits:
        display = clean_text(raw) or "<EMPTY>"
        return display, "not_bitfield", False, [name for name, _exp in expected_pairs]
    parts = []
    missing = []
    all_match = True
    for name, expected in expected_pairs:
        actual = dump_bits.get(norm_key(name))
        if actual is None:
            missing.append(name)
            all_match = False
            parts.append(f"{name}=<MISSING>")
            continue
        parts.append(f"{name}-{actual}")
        if not values_match(actual, expected):
            all_match = False
    state = "bit_missing" if missing else "found"
    return "&".join(parts), state, all_match, missing


def load_mapping(path: Path):
    moc_to_sheet = {}
    attrs_by_moc = defaultdict(dict)
    attrs_global = defaultdict(list)
    rows = read_sheet_rows(path, "MAPPING DEF")
    if not rows:
        raise ValueError(f"MAPPING DEF sheet not found in {path}")
    for row in rows[1:]:
        if not row or len(row) < 5:
            continue
        sheet_name, _group, column, moc, attr = row[0], row[1], row[2], row[3], row[4]
        if not sheet_name or not moc or not attr:
            continue
        moc_key = norm_key(moc)
        attr_key = norm_key(attr)
        moc_to_sheet.setdefault(moc_key, sheet_name)
        moc_to_sheet.setdefault(norm_key(sheet_name), sheet_name)
        rec = {
            "sheet": sheet_name,
            "column": column,
            "moc": moc,
            "attr": attr,
        }
        attrs_by_moc[moc_key][attr_key] = rec
        attrs_global[attr_key].append(rec)
        attrs_global[norm_key(column)].append(rec)
    return {
        "path": path,
        "moc_to_sheet": moc_to_sheet,
        "attrs_by_moc": attrs_by_moc,
        "attrs_global": attrs_global,
    }


class SheetCache:
    def __init__(self, path: Path):
        self.path = path
        self.cache = {}
        self._names = None

    def names(self):
        if self._names is None:
            self._names = sheet_names(self.path)
        return self._names

    def get(self, sheet_name: str):
        if sheet_name in self.cache:
            return self.cache[sheet_name]
        if sheet_name not in self.names():
            self.cache[sheet_name] = None
            return None
        rows = read_sheet_rows(self.path, sheet_name)
        if not rows or len(rows) < 1:
            self.cache[sheet_name] = None
            return None
        huawei_mode = any(norm_key(n) == "MAPPINGDEF" for n in self.names())
        if huawei_mode and len(rows) >= 2:
            header_row, data_start = 1, 2
        else:
            header_row, data_start = detect_data_header_row(rows)
        if header_row >= len(rows):
            self.cache[sheet_name] = None
            return None
        raw_headers = rows[header_row]
        headers = [clean_text(h) if h is not None else f"col_{i}" for i, h in enumerate(raw_headers)]
        header_index = {h: i for i, h in enumerate(headers) if h}
        header_norm = {norm_key(h): i for i, h in enumerate(headers) if norm_key(h)}
        if header_row > 0:
            for i, h in enumerate(rows[0]):
                nk = norm_key(h)
                if nk and nk not in header_norm:
                    header_norm[nk] = i
                    text = clean_text(h)
                    if text and text not in header_index:
                        header_index[text] = i
        records = []
        for row in rows[data_start:]:
            if not row or all(is_empty(v) for v in row):
                continue
            records.append(row)
        payload = {
            "headers": headers,
            "header_index": header_index,
            "header_norm": header_norm,
            "rows": records,
            "identity_idx": smart.identity_indexes(header_norm, fuzzy=False),
        }
        self.cache[sheet_name] = payload
        return payload


def detect_data_header_row(rows) -> tuple[int, int]:
    """Return (header_row_index, first_data_row_index).

    Huawei bulk dumps typically have technical names on row 1 and display
    names on row 2. Generic workbooks use row 1 as headers.
    """
    if len(rows) >= 2:
        first = [norm_key(c) for c in (rows[0] or [])[:12]]
        huawei_markers = {"MODIND", "MOI", "FDN", "NE", "BN", "SN"}
        if any(m in first for m in huawei_markers):
            return 1, 2
        row0_filled = [c for c in (rows[0] or []) if not is_empty(c)]
        row1_filled = [c for c in (rows[1] or []) if not is_empty(c)]
        row1_star = sum(1 for c in row1_filled if str(c).lstrip().startswith("*"))
        if row1_star >= 1 and len(row1_filled) >= len(row0_filled):
            return 1, 2
        if len(row0_filled) <= 3 and len(row1_filled) >= 4:
            return 1, 2
    return 0, 1


def find_attr_record(mapping, moc_name, attr_name):
    moc_key = norm_key(moc_name)
    attr_key = norm_key(attr_name)
    by_moc = mapping["attrs_by_moc"].get(moc_key, {})
    if attr_key in by_moc:
        return by_moc[attr_key]
    for rec in mapping["attrs_global"].get(attr_key, []):
        if norm_key(rec["sheet"]) == moc_key or norm_key(rec["moc"]) == moc_key:
            return rec
    hits = mapping["attrs_global"].get(attr_key, [])
    if len(hits) == 1:
        return hits[0]
    return None


def resolve_parameter(param, mapping):
    mml = clean_text(param["mml"])
    pid = clean_text(param["pid"])
    pname = clean_text(param.get("param_name"))
    result = {
        "sheet": None,
        "column": None,
        "attr": None,
        "bit": None,
        "moc": None,
        "reason": "",
        "smart_note": "",
    }
    if not mml or (not pid and not pname):
        result["reason"] = "Missing MML Object, or both Parameter ID and Parameter Name"
        return result

    notes = []
    moc_names = []
    for recs in mapping["attrs_by_moc"].values():
        for rec in recs.values():
            moc_names.append(rec["moc"])
            moc_names.append(rec["sheet"])
    moc_names.extend(mapping["moc_to_sheet"].values())
    matched_mo, score, why = smart.closest_name(mml, moc_names, cutoff=0.8)
    if matched_mo and smart.norm_key(matched_mo) != smart.norm_key(mml):
        notes.append(f"Interpreted MO '{mml}' as '{matched_mo}' ({why})")
        mml = matched_mo

    moc_key = norm_key(mml)
    moc_attrs = mapping["attrs_by_moc"].get(moc_key, {})
    mml_as_attr_hits = mapping["attrs_global"].get(moc_key, [])

    attr_name, bit_name = smart.split_parameter_id(pid)
    if "@" not in pid:
        attr_name, bit_name = pid, None

    def attr_names_for_mo():
        names = []
        for rec in moc_attrs.values():
            names.append(rec["attr"])
            names.append(rec["column"])
        return names

    def locate(attr):
        if not attr:
            return None
        rec = find_attr_record(mapping, mml, attr)
        if rec:
            return rec
        hit, hit_score, hit_why = smart.closest_name(attr, attr_names_for_mo(), cutoff=0.8)
        if hit:
            rec = find_attr_record(mapping, mml, hit)
            if rec:
                notes.append(f"Interpreted parameter '{attr}' as '{hit}' ({hit_why})")
            return rec
        return None

    rec = None
    if pname:
        rec = locate(pname)
        if rec:
            notes.append(f"Matched dump using Parameter Name '{pname}'")
    if rec is None and attr_name:
        rec = locate(attr_name)
    if rec is None and bit_name:
        swapped = locate(bit_name)
        if swapped:
            rec = swapped
            attr_name, bit_name = bit_name, attr_name
            notes.append("Parameter ID sides around @ were swapped to match the dump")
    if rec is None and "@" not in pid:
        if moc_attrs and norm_key(pid) in moc_attrs:
            rec = moc_attrs[norm_key(pid)]
        elif mml_as_attr_hits and not moc_attrs:
            rec = mml_as_attr_hits[0]
            if bit_name is None:
                bit_name = pid
            attr_name = rec["attr"]
    if rec is None and mml_as_attr_hits:
        rec = mml_as_attr_hits[0]
        if bit_name is None:
            bit_name = pid if "@" not in pid else smart.split_parameter_id(pid)[1]
        attr_name = rec["attr"]
    if rec is None:
        result["reason"] = f"Attribute not mapped: {mml} / {pid}"
        if not pname:
            result["reason"] += ". If this sheet has no Parameter Name, add Parameter Name along with Parameter ID."
        result["smart_note"] = " | ".join(notes)
        return result

    result.update(
        {
            "sheet": rec["sheet"],
            "column": rec["column"],
            "attr": rec["attr"],
            "bit": bit_name,
            "moc": rec["moc"],
            "reason": "OK",
            "smart_note": " | ".join(notes),
        }
    )
    return result


def object_label(row, headers, ctx=None):
    if ctx and ctx.cell_key:
        extra = []
        if ctx.band_family or ctx.band_raw:
            extra.append(ctx.band_family or ctx.band_raw)
        label = ctx.cell_key
        if extra:
            label += f" [{', '.join(extra)}]"
        return label
    parts = []
    for key in headers[:3]:
        idx = headers.index(key)
        val = row[idx] if idx < len(row) else None
        if not is_empty(val):
            parts.append(f"{key}={val}")
    return " | ".join(parts[:2]) if parts else "row"


class CellBandIndex:
    """Map Local cell ID → L09/L18/L21/L26 from Cell (4G) and NRDUCELL (5G) only.

    Other MO sheets are not loaded. Group IDs are read from the current
    parameter row when that sheet has them.
    """

    def __init__(self, cache: SheetCache):
        self.by_key = {}
        self.by_cell_id = {}
        self._loaded_sheets = []
        self._build(cache)

    def _sheet_role(self, sheet_name: str) -> str | None:
        nk = norm_key(sheet_name)
        cell_names = {"CELL", *{norm_key(a) for a in smart.SHEET_ALIASES["CELL"]}}
        nr_names = {"NRDUCELL", *{norm_key(a) for a in smart.SHEET_ALIASES["NRDUCELL"]}}
        if nk in cell_names:
            return "cell"
        if nk in nr_names:
            return "nrducell"
        return None

    def _build(self, cache: SheetCache):
        for sheet_name in cache.names():
            role = self._sheet_role(sheet_name)
            if role not in {"cell", "nrducell"}:
                continue
            payload = cache.get(sheet_name)
            if not payload:
                continue
            self._loaded_sheets.append(sheet_name)
            for row in payload["rows"]:
                ctx = smart.cell_from_row(
                    payload["headers"], payload["header_norm"], row, indexes=payload.get("identity_idx")
                )
                if not ctx.cell_id:
                    continue
                cid = norm_key(ctx.cell_id)
                if ctx.cell_key:
                    existing = self.by_key.get(ctx.cell_key)
                    if existing is None:
                        self.by_key[ctx.cell_key] = ctx
                    else:
                        existing.band_tokens |= ctx.band_tokens
                        if not existing.band_family:
                            existing.band_family = ctx.band_family
                        if not existing.band_raw:
                            existing.band_raw = ctx.band_raw
                if cid and cid not in self.by_cell_id:
                    self.by_cell_id[cid] = ctx

    def enrich(self, ctx: smart.RowContext) -> smart.RowContext:
        indexed = None
        if ctx.cell_key:
            indexed = self.by_key.get(ctx.cell_key)
        if indexed is None and ctx.cell_id:
            indexed = self.by_cell_id.get(norm_key(ctx.cell_id))
        if indexed:
            if not ctx.band_tokens:
                ctx.band_tokens = set(indexed.band_tokens)
                ctx.band_family = indexed.band_family
                ctx.band_raw = indexed.band_raw
            if not ctx.cell_name:
                ctx.cell_name = indexed.cell_name
        return ctx


def compare_param(param, resolved, cache, cell_index=None):
    recommend = param["recommend"]
    spec = smart.parse_recommend(recommend)
    out = {
        "status": "",
        "objects_checked": 0,
        "match_count": 0,
        "mismatch_count": 0,
        "missing_count": 0,
        "skipped_not_applicable": 0,
        "unique_actuals": [],
        "sample_mismatches": [],
        "remark": resolved.get("reason") or "",
        "actual_source": "",
        "applied_recommend": "",
    }
    if resolved.get("smart_note"):
        out["remark"] = ((out["remark"] + " | ") if out["remark"] and out["remark"] != "OK" else "") + resolved["smart_note"]
        if resolved.get("reason") == "OK" and not out["remark"]:
            out["remark"] = resolved["smart_note"]
    if param.get("no_recommend_reason") and is_empty(recommend):
        out["status"] = "No Recommend Value"
        out["remark"] = param["no_recommend_reason"]
        return out
    if resolved.get("reason") != "OK":
        out["status"] = "Not Found in Configuration"
        return out

    sheet = cache.get(resolved["sheet"])
    if not sheet:
        out["status"] = "Not Found in Configuration"
        out["remark"] = f"Sheet not present: {resolved['sheet']}"
        return out

    col = resolved["column"]
    col_idx = find_column_index(sheet, col)
    if col_idx is None:
        pname = clean_text(param.get("param_name"))
        if pname:
            col_idx = find_column_index(sheet, pname)
            if col_idx is None:
                hit, _score, why = smart.closest_name(
                    pname, sheet["headers"], cutoff=0.8
                )
                if hit:
                    col_idx = find_column_index(sheet, hit) or sheet["header_index"].get(hit)
                    out["remark"] = (
                        (out["remark"] + " | " if out["remark"] and out["remark"] != "OK" else "")
                        + f"Interpreted column '{pname}' as '{hit}' ({why})"
                    )
    if col_idx is None:
        out["status"] = "Not Found in Configuration"
        out["remark"] = f"Column not present: {col}"
        if not param.get("param_name"):
            out["remark"] += ". Add Parameter Name along with Parameter ID if this field is a display name."
        return out

    out["actual_source"] = f"{resolved['sheet']} / {sheet['headers'][col_idx]}"
    counts = Counter()
    mismatches = []
    applied_notes = Counter()
    cond_clauses = []
    for text in condition_texts(param):
        cond_clauses.extend(smart.parse_conditions(text))
    need_ctx = spec.has_conditional
    indexes = sheet.get("identity_idx") if need_ctx else None
    bit_name = resolved.get("bit")
    rec_switch_bits = None if need_ctx else parse_switch_recommend(recommend)
    if rec_switch_bits and not bit_name:
        resolved["bit"] = "&".join(name for name, _exp in rec_switch_bits)
        bit_name = resolved["bit"]
    cond_cols = []
    if cond_clauses:
        cond_cols, missing_cond = resolve_condition_columns(sheet, cond_clauses)
        if missing_cond:
            out["status"] = "Not Found in Configuration"
            out["remark"] = (
                "Conditions refers to dump column(s) not found on "
                f"{resolved['sheet']}: {', '.join(missing_cond)}. "
                "Use the dump Parameter Name and short name in brackets, "
                "e.g. Interfreq handover group ID (INTERFREQHOGROUPID)=0"
            )
            return out
    cond_suffix = ""
    if cond_cols:
        cond_suffix = " [" + ", ".join(
            f"{clause.display_name}={clause.value}" for _idx, _val, clause in cond_cols
        ) + "]"
    for row in sheet["rows"]:
        if cond_cols and not row_matches_conditions(row, cond_cols):
            out["skipped_not_applicable"] += 1
            continue
        raw = row[col_idx] if col_idx < len(row) else None
        ctx = None
        if need_ctx:
            ctx = smart.cell_from_row(sheet["headers"], sheet["header_norm"], row, indexes=indexes)
            if cell_index is not None:
                ctx = cell_index.enrich(ctx)
            expected, applied_note = smart.select_recommend(spec, ctx)
            if expected is None:
                out["skipped_not_applicable"] += 1
                continue
            row_recommend = expected
            if applied_note:
                applied_notes[f"{applied_note}=>{row_recommend}"] += 1
        else:
            row_recommend = recommend

        switch_bits = parse_switch_recommend(row_recommend) if need_ctx else rec_switch_bits
        used = raw
        bit_state = ""
        switch_all_match = None
        if switch_bits:
            used, bit_state, switch_all_match, missing_bits = audit_switch_bits(raw, switch_bits)
            if missing_bits and "closest_bit" not in out:
                dump_bits = parse_bitfield(raw) or {}
                hit, _score, _why = smart.closest_name(missing_bits[0], dump_bits.keys(), cutoff=0.8)
                if hit:
                    out["closest_bit"] = hit
            if not resolved.get("bit"):
                resolved["bit"] = "&".join(name for name, _exp in switch_bits)
        elif bit_name:
            bit_val, bit_state = extract_bit(raw, bit_name)
            if bit_state == "found":
                used = bit_val
            elif bit_state == "not_bitfield":
                used = raw
            else:
                used = f"<BIT_MISSING:{bit_name}>"
                if "closest_bit" not in out:
                    bits = parse_bitfield(raw) or {}
                    similar = [
                        name
                        for name in bits
                        if bit_name.replace("NSA_", "NR_") in name
                        or name.endswith(norm_key(bit_name)[-16:])
                    ]
                    if not similar:
                        hit, _score, _why = smart.closest_name(bit_name, bits.keys(), cutoff=0.8)
                        if hit:
                            similar = [hit]
                    if similar:
                        out["closest_bit"] = similar[0]
        display = clean_text(used)
        if display == "":
            display = "<EMPTY>"
        counts[display] += 1
        out["objects_checked"] += 1
        if is_empty(row_recommend):
            continue
        if switch_all_match is True:
            out["match_count"] += 1
            continue
        if switch_all_match is False:
            if bit_state == "bit_missing":
                out["missing_count"] += 1
            out["mismatch_count"] += 1
            if len(mismatches) < 8:
                mismatches.append(
                    object_label(row, sheet["headers"], ctx)
                    + cond_suffix
                    + f" expected {clean_text(row_recommend)} -> {display}"
                )
            continue
        if bit_state == "bit_missing":
            out["missing_count"] += 1
            out["mismatch_count"] += 1
            if len(mismatches) < 8:
                mismatches.append(
                    object_label(row, sheet["headers"], ctx) + cond_suffix + f" -> {display}"
                )
            continue
        if values_match(used, row_recommend):
            out["match_count"] += 1
        else:
            out["mismatch_count"] += 1
            if len(mismatches) < 8:
                mismatches.append(
                    object_label(row, sheet["headers"], ctx)
                    + cond_suffix
                    + f" expected {clean_text(row_recommend)} -> {display}"
                )

    out["unique_actuals"] = counts.most_common(12)
    out["sample_mismatches"] = mismatches
    if applied_notes:
        out["applied_recommend"] = " | ".join(f"{k} ({n})" for k, n in applied_notes.most_common(8))
    if cond_suffix:
        note = cond_suffix.strip()
        out["applied_recommend"] = (
            f"{note} | {out['applied_recommend']}" if out["applied_recommend"] else note
        )
    if param.get("no_recommend_reason") and is_empty(recommend):
        out["status"] = "No Recommend Value"
        out["remark"] = param["no_recommend_reason"]
        return out
    if is_empty(recommend):
        out["status"] = "No Recommend Value"
        out["remark"] = "Reference recommend value is empty; parameter not audited for inconsistency"
        return out
    if out["objects_checked"] == 0:
        if out["skipped_not_applicable"]:
            out["status"] = "No Recommend Value"
            out["remark"] = conditions_skip_remark(param, out["skipped_not_applicable"])
            return out
        out["status"] = "Not Found in Configuration"
        out["remark"] = "No data rows in mapped sheet"
        return out
    if out["mismatch_count"] == 0:
        out["status"] = "Consistent"
    elif out["match_count"] == 0:
        out["status"] = "Inconsistent"
    else:
        out["status"] = "Mixed / Partial"

    pid = str(param.get("pid") or "").upper()
    if pid == "DLARFCN" and out["status"] == "Inconsistent":
        out["remark"] = (
            "Recommend is a band name, but DlArfcn in the dump is a numeric ARFCN. "
            f"Actual: {unique_actual_text(out['unique_actuals'])}. "
            "Check Frequency Band on the same MO."
        )
    elif pid == "FREQUENCYBAND" and out["status"] == "Inconsistent":
        out["remark"] = (
            "Recommend does not match Frequency Band in the dump. "
            f"Actual: {unique_actual_text(out['unique_actuals'])}. "
            "Additional Frequency Band may be the intended NULL field."
        )
    elif out["unique_actuals"] and (
        str(out["unique_actuals"][0][0]).startswith("<BIT_MISSING")
        or "=<MISSING>" in str(out["unique_actuals"][0][0])
    ):
        extra = ""
        if "closest_bit" in out:
            extra = f" Closest live bit key seen: {out['closest_bit']}."
        out["remark"] = (
            f"Bit {resolved.get('bit')} is not present in the packed switch on "
            f"{resolved.get('sheet')} / {resolved.get('column')} for this software version."
            + extra
        )
    if resolved.get("smart_note") and out["remark"] in {"", "OK"}:
        out["remark"] = resolved["smart_note"]
    elif resolved.get("smart_note") and resolved["smart_note"] not in (out["remark"] or ""):
        out["remark"] = (out["remark"] + " | " if out["remark"] else "") + resolved["smart_note"]
    return out


IDENTITY_REF_FIELDS = ("mml", "pid", "param_name")
HEADER_PROBES = {
    "mml": ("MML Object", "MO Name", "MML Object Name", "Managed Object", "MOC"),
    "pid": ("Parameter ID", "Param ID", "Para ID", "ParameterId"),
    "param_name": ("Parameter Name", "Param Name", "Display Name", "Full Name"),
    "recommend": (
        "Proposed Value",
        "Recommended Value",
        "Recommend Value",
        "Plan Value",
        "Target Value",
        "Golden Value",
        "Expected Value",
    ),
}
SKIP_AS_RECOMMEND = {
    "DEFAULTVALUE",
    "DEFAULT",
    "FUNCTION",
    "FEATURE",
    "REMARK",
    "REMARKS",
    "COMMENT",
    "COMMENTS",
    "NOTE",
    "NOTES",
    "UNIT",
    "UNITS",
    "DESCRIPTION",
    "RANGE",
    "TARGETIMPACT",
    "TARGIMPACT",
    "IMPACT",
}


def detect_ref_header_map(header_row) -> dict:
    """Map identity columns plus adaptive recommend and Conditions1..N columns.

    Identity (always expected): MML Object, Parameter ID, Parameter Name.
    Recommend / Proposed / Plan is found adaptively. Conditions1, Conditions2,
    … ConditionsN (or unnumbered Conditions) are detected by name and are
    never treated as the recommended value.
    """
    mapping = {}
    used = set()
    cells = []
    for i, cell in enumerate(header_row or []):
        text = clean_text(cell)
        key = norm_key(cell)
        cells.append((i, text, key))
        if not key:
            continue
        if smart.condition_header_slot(text) is not None:
            continue
        for field, aliases in REF_HEADER_ALIASES.items():
            if key in aliases and field not in mapping:
                mapping[field] = i
                used.add(i)
                break
    unused = [(i, text) for i, text, key in cells if i not in used and text]
    for field, probes in HEADER_PROBES.items():
        if field in mapping:
            continue
        best_i = None
        best_score = 0.0
        for i, text in unused:
            if field == "recommend" and (
                norm_key(text) in SKIP_AS_RECOMMEND
                or smart.condition_header_slot(text) is not None
            ):
                continue
            for probe in probes:
                score = smart.similarity(text, probe)
                if score > best_score:
                    best_i, best_score = i, score
        if best_i is not None and best_score >= 0.78:
            mapping[field] = best_i
            used.add(best_i)
            unused = [(i, text) for i, text in unused if i != best_i]
    mapping["condition_cols"] = collect_condition_columns(cells, used)
    if 1 in mapping["condition_cols"]:
        mapping["conditions"] = mapping["condition_cols"][1]
    return mapping


def collect_condition_columns(cells, used: set) -> dict:
    """Map Conditions1..N (and unnumbered Conditions) to column indexes."""
    numbered = {}
    unnumbered = []
    for i, text, key in cells:
        if i in used or not key:
            continue
        slot = smart.condition_header_slot(text)
        if slot is None:
            continue
        if smart.condition_header_is_numbered(text):
            if slot not in numbered:
                numbered[slot] = i
        else:
            unnumbered.append((i, slot))
    for i, slot in unnumbered:
        if slot not in numbered:
            numbered[slot] = i
    for i in numbered.values():
        used.add(i)
    return numbered


def is_identity_ref_header(mapping: dict) -> bool:
    return all(field in mapping for field in IDENTITY_REF_FIELDS)


def is_usable_ref_header(mapping: dict) -> bool:
    if is_identity_ref_header(mapping):
        return True
    return "pid" in mapping and "mml" in mapping and "recommend" in mapping


def is_skipped_ref_sheet(sheet_name: str) -> bool:
    key = norm_key(sheet_name)
    skipped = {norm_key(s) for s in SKIP_REF_SHEETS | META_SHEETS}
    return key in skipped


def guess_network_label(*parts) -> str:
    blob = " ".join(clean_text(p) for p in parts if p).upper()
    if re.search(r"(^|[^0-9])5G([^0-9]|$)", blob) or "NR PERFORMANCE" in blob or "GNODEB" in blob:
        return "5G"
    if re.search(r"(^|[^0-9])4G([^0-9]|$)", blob) or "NR ANCHOR" in blob or re.search(r"\bLTE\b", blob):
        return "4G"
    if re.search(r"(^|[^0-9])2G([^0-9]|$)", blob) or "GSM" in blob or "BSC" in blob:
        return "2G"
    if "3G" in blob or "NODEB" in blob or "RNC" in blob or "UMTS" in blob:
        return "3G"
    return ""


def load_reference_file(ref_path: Path, ref_count: int = 1):
    params = []
    names = sheet_names(ref_path)
    for sheet_name in names:
        if is_skipped_ref_sheet(sheet_name):
            continue
        rows = read_sheet_rows(ref_path, sheet_name)
        if not rows:
            continue
        header_idx = None
        header_map = None
        best_score = -1
        for i, row in enumerate(rows[:12]):
            candidate = detect_ref_header_map(row)
            if not is_usable_ref_header(candidate):
                continue
            score = (
                (3 if "mml" in candidate else 0)
                + (3 if "pid" in candidate else 0)
                + (3 if "param_name" in candidate else 0)
                + (1 if "recommend" in candidate else 0)
                + (1 if candidate.get("condition_cols") else 0)
            )
            if score > best_score:
                best_score = score
                header_idx = i
                header_map = candidate
        if header_map is None:
            continue
        sheet_as_mo = "mml" not in header_map
        missing_name_col = "param_name" not in header_map
        for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
            def cell(field):
                idx = header_map.get(field)
                if idx is None or row is None or idx >= len(row):
                    return None
                return row[idx]

            mml = clean_text(cell("mml"))
            pid = clean_text(cell("pid"))
            pname = clean_text(cell("param_name"))
            if is_empty(mml) and is_empty(pid) and is_empty(pname):
                continue
            if is_empty(pid) and is_empty(pname):
                continue
            if is_empty(mml):
                mml = sheet_name
                sheet_as_mo = True
            rec = cell("recommend")
            no_recommend_reason = None
            new_style = smart.is_new_style_ref_headers(rows[header_idx])
            if missing_name_col and new_style:
                rec = None
                no_recommend_reason = (
                    "No Parameter Name column. Add MML Object, Parameter ID, and Parameter Name."
                )
            elif missing_name_col is False and is_empty(pname):
                rec = None
                no_recommend_reason = (
                    "Parameter Name is empty. Add Parameter Name along with MML Object and Parameter ID."
                )
            sheet_key = sheet_name if ref_count <= 1 else f"{ref_path.name} | {sheet_name}"
            cond_cols = header_map.get("condition_cols") or {}
            condition_values = {}
            for n, idx in sorted(cond_cols.items()):
                raw = row[idx] if row is not None and idx < len(row) else None
                txt = clean_text(raw)
                condition_values[int(n)] = txt or None
            params.append(
                {
                    "ref_file": ref_path.name,
                    "ref_path": str(ref_path),
                    "ref_sheet": sheet_name,
                    "sheet_key": sheet_key,
                    "ref_row": offset,
                    "function": clean_text(cell("function")) or "Unspecified",
                    "mml": mml,
                    "pid": pid,
                    "param_name": pname or None,
                    "recommend": rec if not is_empty(rec) else None,
                    "conditions": condition_values.get(1),
                    "condition_values": condition_values,
                    "condition_slots": sorted(cond_cols),
                    "no_recommend_reason": no_recommend_reason,
                    "network": guess_network_label(ref_path.name, sheet_name),
                    "sheet_is_mo": sheet_as_mo,
                }
            )
    return params


def load_reference(ref_path: Path):
    """Load one reference workbook (any qualifying sheets, not only NR Performance / NR Anchor)."""
    return load_reference_file(Path(ref_path), ref_count=1)


def load_all_references(ref_files: list[Path]):
    params = []
    missing = []
    for path in ref_files:
        loaded = load_reference_file(path, ref_count=len(ref_files))
        if not loaded:
            missing.append(path.name)
            continue
        params.extend(loaded)
    return params, missing


def find_column_index(sheet, name: str):
    if not name or not sheet:
        return None
    idx = sheet["header_index"].get(name)
    if idx is not None:
        return idx
    idx = sheet["header_norm"].get(norm_key(name))
    if idx is not None:
        return idx
    wanted = smart.header_match_key(name)
    if wanted:
        for header in sheet.get("headers") or []:
            if smart.header_match_key(header) == wanted:
                return sheet["header_index"].get(header)
    return None


def resolve_condition_columns(sheet, clauses):
    """Map each Conditions clause to a dump column index (once per parameter)."""
    resolved = []
    missing = []
    identity_groups = (sheet.get("identity_idx") or {}).get("groups") or []
    group_by_key = {norm_key(key): idx for key, idx in identity_groups}
    header_norm = sheet.get("header_norm") or {}
    for clause in clauses:
        idx = None
        for label in clause.labels():
            idx = find_column_index(sheet, label)
            if idx is not None:
                break
        if idx is None:
            for key in smart.condition_name_keys(clause):
                if key in header_norm:
                    idx = header_norm[key]
                    break
                if key in group_by_key:
                    idx = group_by_key[key]
                    break
        if idx is None:
            candidates = list(sheet.get("headers") or [])
            for label in clause.labels():
                hit, _score, _why = smart.closest_name(label, candidates, cutoff=0.82)
                if hit:
                    idx = find_column_index(sheet, hit)
                    if idx is not None:
                        break
        if idx is None:
            missing.append(clause.display_name or clause.short_name or clause.raw)
        else:
            resolved.append((idx, clause.value, clause))
    return resolved, missing


def row_matches_conditions(row, cond_cols) -> bool:
    for idx, expected, _clause in cond_cols:
        actual = row[idx] if row is not None and idx < len(row) else None
        if smart.canon_id_value(actual) != expected:
            return False
    return True


def condition_texts(param) -> list[str]:
    values = param.get("condition_values") or {}
    texts = [values[n] for n in sorted(values) if not is_empty(values.get(n))]
    if texts:
        return texts
    if param.get("conditions"):
        return [param["conditions"]]
    return []


def condition_slots_for_report(params) -> list[int]:
    slots = set()
    for p in params or []:
        slots.update(p.get("condition_slots") or [])
        slots.update((p.get("condition_values") or {}).keys())
        if p.get("conditions") and 1 not in slots:
            slots.add(1)
    return sorted(int(n) for n in slots if str(n).isdigit() or isinstance(n, int))


def condition_value_for_slot(param, slot: int) -> str:
    values = param.get("condition_values") or {}
    if slot in values and values[slot] is not None:
        return values[slot]
    if int(slot) == 1:
        return param.get("conditions") or ""
    return ""


def conditions_skip_remark(param, skipped: int) -> str:
    bits = []
    values = param.get("condition_values") or {}
    if values:
        for n in sorted(values):
            if not is_empty(values[n]):
                bits.append(f"Conditions{n}={values[n]}")
    elif param.get("conditions"):
        bits.append(f"Conditions1={param['conditions']}")
    if bits:
        return (
            "No cell matched " + "; ".join(bits) + f"; {skipped} row(s) skipped. "
            "Objects checked is the count of rows where every Conditions column matches."
        )
    return (
        "No cell matched the band/group conditions in Proposed/Recommended Value "
        f"({skipped} row(s) not applicable)."
    )


def resolve_by_sheet_name(param, workbook_names, cache):
    """Treat a matching input sheet name as the MML / MO object."""
    result = {
        "sheet": None,
        "column": None,
        "attr": None,
        "bit": None,
        "moc": None,
        "reason": "",
        "smart_note": "",
    }
    mml = clean_text(param.get("mml"))
    pid = clean_text(param.get("pid"))
    pname = clean_text(param.get("param_name"))
    notes = []
    if not pid and not pname:
        result["reason"] = "Missing Parameter ID and Parameter Name"
        return result
    name_index = {norm_key(n): n for n in workbook_names if not is_skipped_ref_sheet(n) and norm_key(n) not in {norm_key(s) for s in META_SHEETS}}
    sheet = name_index.get(norm_key(mml)) if mml else None
    if sheet is None and mml:
        for key, actual in name_index.items():
            if key.endswith(norm_key(mml)) or norm_key(mml).endswith(key):
                if min(len(key), len(norm_key(mml))) >= 6:
                    sheet = actual
                    notes.append(f"Interpreted MO '{mml}' as sheet '{actual}'")
                    break
    if sheet is None and mml:
        hit, _score, why = smart.closest_name(mml, name_index.values(), cutoff=0.8)
        if hit:
            sheet = hit
            notes.append(f"Interpreted MO '{mml}' as sheet '{hit}' ({why})")
    if sheet is None:
        result["reason"] = f"No input sheet named as MO '{mml}'"
        result["smart_note"] = " | ".join(notes)
        return result
    payload = cache.get(sheet)
    if not payload:
        result["reason"] = f"Sheet not present: {sheet}"
        return result
    attr_name, bit_name = smart.split_parameter_id(pid)
    candidates = [c for c in (pname, attr_name, bit_name, pid) if c]
    col_idx = None
    col_name = None
    chosen = None
    for cand in candidates:
        idx = find_column_index(payload, cand)
        if idx is not None:
            col_idx = idx
            col_name = payload["headers"][idx]
            chosen = cand
            break
    if col_idx is None:
        for cand in candidates:
            hit, _score, why = smart.closest_name(cand, payload["headers"], cutoff=0.8)
            if hit:
                col_idx = find_column_index(payload, hit)
                col_name = hit
                chosen = cand
                notes.append(f"Interpreted column '{cand}' as '{hit}' ({why})")
                break
    if col_idx is None:
        result["reason"] = f"Column not present on sheet {sheet}: {pid}"
        if not pname:
            result["reason"] += ". Add Parameter Name along with Parameter ID."
        result["sheet"] = sheet
        result["smart_note"] = " | ".join(notes)
        return result
    if bit_name is None and attr_name and chosen and smart.norm_key(chosen) == smart.norm_key(attr_name):
        bit_name = None
    if "@" in pid and chosen and smart.norm_key(chosen) == smart.norm_key(attr_name or ""):
        pass
    elif "@" in pid and chosen and smart.norm_key(chosen) == smart.norm_key(bit_name or ""):
        attr_name, bit_name = bit_name, attr_name
    result.update(
        {
            "sheet": sheet,
            "column": col_name,
            "attr": col_name,
            "bit": bit_name if chosen and smart.norm_key(chosen) != smart.norm_key(bit_name or "") else (
                bit_name if "@" in pid else None
            ),
            "moc": mml or sheet,
            "reason": "OK",
            "smart_note": " | ".join(notes),
        }
    )
    if "@" in pid:
        _attr, switch_name = smart.split_parameter_id(pid)
        result["bit"] = switch_name
    return result


class InputWorkbook:
    def __init__(self, path: Path):
        self.path = Path(path)
        self.cache = SheetCache(self.path)
        self._mapping = None
        self._mapping_loaded = False
        self._cell_index = None

    @property
    def mapping(self):
        if not self._mapping_loaded:
            self._mapping_loaded = True
            try:
                names = self.cache.names()
            except Exception:
                names = []
            if any(norm_key(n) == "MAPPINGDEF" for n in names):
                try:
                    self._mapping = load_mapping(self.path)
                except Exception:
                    self._mapping = None
            else:
                self._mapping = None
        return self._mapping

    def cell_index(self) -> CellBandIndex:
        if self._cell_index is None:
            self._cell_index = CellBandIndex(self.cache)
        return self._cell_index

    def resolve(self, param) -> dict:
        moc_key = norm_key(param.get("mml"))
        if self.mapping:
            mo_known = (
                moc_key in self.mapping["moc_to_sheet"]
                or moc_key in self.mapping["attrs_by_moc"]
                or moc_key in self.mapping["attrs_global"]
            )
            if not mo_known and param.get("mml"):
                moc_names = list(self.mapping["moc_to_sheet"].values())
                for recs in self.mapping["attrs_by_moc"].values():
                    for rec in recs.values():
                        moc_names.append(rec["moc"])
                        moc_names.append(rec["sheet"])
                hit, _score, _why = smart.closest_name(param.get("mml"), moc_names, cutoff=0.8)
                if hit:
                    mo_known = True
            if mo_known:
                resolved = resolve_parameter(param, self.mapping)
                if resolved.get("reason") == "OK":
                    return resolved
        return resolve_by_sheet_name(param, self.cache.names(), self.cache)


class InputStore:
    def __init__(self, files: list[Path]):
        self.workbooks = [InputWorkbook(path) for path in files]

    def compare_param(self, param):
        if param.get("no_recommend_reason") and is_empty(param.get("recommend")):
            resolved = {
                "sheet": None,
                "column": None,
                "attr": None,
                "bit": None,
                "moc": None,
                "reason": param["no_recommend_reason"],
                "smart_note": "",
            }
            empty = {
                "status": "No Recommend Value",
                "objects_checked": 0,
                "match_count": 0,
                "mismatch_count": 0,
                "missing_count": 0,
                "skipped_not_applicable": 0,
                "unique_actuals": [],
                "sample_mismatches": [],
                "remark": param["no_recommend_reason"],
                "actual_source": "",
                "input_files": "",
                "applied_recommend": "",
            }
            return resolved, empty
        hits = []
        last_resolved = {
            "sheet": None,
            "column": None,
            "attr": None,
            "bit": None,
            "moc": None,
            "reason": "Not found in any input file",
            "smart_note": "",
        }
        for wb in self.workbooks:
            resolved = wb.resolve(param)
            if resolved.get("reason") != "OK":
                last_resolved = resolved
                continue
            result = compare_param(
                param,
                resolved,
                wb.cache,
                cell_index=wb.cell_index() if smart.parse_recommend(param.get("recommend")).has_conditional else None,
            )
            if result["status"] == "Not Found in Configuration" and result["objects_checked"] == 0:
                last_resolved = resolved
                last_resolved["reason"] = result.get("remark") or resolved.get("reason")
                continue
            result["input_file"] = wb.path.name
            result["resolved"] = resolved
            hits.append((wb.path, resolved, result))
        if not hits:
            empty = {
                "status": "Not Found in Configuration",
                "objects_checked": 0,
                "match_count": 0,
                "mismatch_count": 0,
                "missing_count": 0,
                "skipped_not_applicable": 0,
                "unique_actuals": [],
                "sample_mismatches": [],
                "remark": last_resolved.get("reason") or "Not found in any input file",
                "actual_source": "",
                "input_files": "",
                "applied_recommend": "",
            }
            return last_resolved, empty
        return merge_file_hits(param, hits)


def merge_file_hits(param, hits):
    primary_resolved = hits[0][1]
    merged = {
        "status": "",
        "objects_checked": 0,
        "match_count": 0,
        "mismatch_count": 0,
        "missing_count": 0,
        "skipped_not_applicable": 0,
        "unique_actuals": [],
        "sample_mismatches": [],
        "remark": "",
        "actual_source": "",
        "applied_recommend": "",
        "input_files": ", ".join(path.name for path, _, _ in hits),
    }
    counts = Counter()
    remarks = []
    sources = []
    applied = []
    for path, resolved, result in hits:
        merged["objects_checked"] += result["objects_checked"]
        merged["match_count"] += result["match_count"]
        merged["mismatch_count"] += result["mismatch_count"]
        merged["missing_count"] += result.get("missing_count") or 0
        merged["skipped_not_applicable"] += result.get("skipped_not_applicable") or 0
        for value, count in result.get("unique_actuals") or []:
            counts[f"{path.name}: {value}"] += count
        for sample in result.get("sample_mismatches") or []:
            if len(merged["sample_mismatches"]) < 12:
                merged["sample_mismatches"].append(f"{path.name} | {sample}")
        if result.get("remark"):
            remarks.append(f"{path.name}: {result['remark']}")
        if result.get("actual_source"):
            sources.append(f"{path.name}: {result['actual_source']}")
        if result.get("applied_recommend"):
            applied.append(f"{path.name}: {result['applied_recommend']}")
        if result.get("closest_bit") and "closest_bit" not in merged:
            merged["closest_bit"] = result["closest_bit"]
    merged["unique_actuals"] = counts.most_common(12)
    merged["remark"] = " | ".join(remarks)
    merged["actual_source"] = " | ".join(sources)
    merged["applied_recommend"] = " | ".join(applied)
    recommend = param.get("recommend")
    if param.get("no_recommend_reason") and is_empty(recommend):
        merged["status"] = "No Recommend Value"
        merged["remark"] = param["no_recommend_reason"]
        return primary_resolved, merged
    if is_empty(recommend):
        merged["status"] = "No Recommend Value"
        merged["remark"] = merged["remark"] or (
            "Reference recommend value is empty; parameter not audited for inconsistency"
        )
        return primary_resolved, merged
    if merged["objects_checked"] == 0:
        if merged["skipped_not_applicable"]:
            merged["status"] = "No Recommend Value"
            merged["remark"] = merged["remark"] or conditions_skip_remark(
                param, merged["skipped_not_applicable"]
            )
            return primary_resolved, merged
        merged["status"] = "Not Found in Configuration"
        return primary_resolved, merged
    if merged["mismatch_count"] == 0:
        merged["status"] = "Consistent"
    elif merged["match_count"] == 0:
        merged["status"] = "Inconsistent"
    else:
        merged["status"] = "Mixed / Partial"
    pid = str(param.get("pid") or "").upper()
    if pid == "DLARFCN" and merged["status"] == "Inconsistent":
        merged["remark"] = (
            "Recommend is a band name, but DlArfcn in the dump is a numeric ARFCN. "
            f"Actual: {unique_actual_text(merged['unique_actuals'])}. "
            "Check Frequency Band on the same MO."
        )
    elif pid == "FREQUENCYBAND" and merged["status"] == "Inconsistent":
        merged["remark"] = (
            "Recommend does not match Frequency Band in the dump. "
            f"Actual: {unique_actual_text(merged['unique_actuals'])}. "
            "Additional Frequency Band may be the intended NULL field."
        )
    elif merged["unique_actuals"] and str(merged["unique_actuals"][0][0]).split(": ", 1)[-1].startswith("<BIT_MISSING"):
        extra = ""
        if "closest_bit" in merged:
            extra = f" Closest live bit key seen: {merged['closest_bit']}."
        merged["remark"] = (
            f"Bit {primary_resolved.get('bit')} is not present in the packed switch on "
            f"{primary_resolved.get('sheet')} / {primary_resolved.get('column')} for this software version."
            + extra
        )
    return primary_resolved, merged


def unique_actual_text(unique_actuals):
    if not unique_actuals:
        return ""
    parts = []
    for value, count in unique_actuals:
        text = str(value)
        if len(text) > 90:
            text = text[:87] + "..."
        parts.append(f"{text} ({count})")
    return " | ".join(parts)


def classify_bucket(status: str) -> str:
    if status == "Inconsistent":
        return "full_inconsistent"
    if status == "Mixed / Partial":
        return "mixed"
    if status == "Consistent":
        return "consistent"
    if status == "No Recommend Value":
        return "no_recommend"
    return "not_found"


def is_inconsistent(status: str) -> bool:
    return classify_bucket(status) in {"full_inconsistent", "mixed"}


def summarize(params):
    summary = defaultdict(lambda: Counter())
    func_summary = defaultdict(lambda: defaultdict(lambda: Counter()))
    for p in params:
        bucket = classify_bucket(p["result"]["status"])
        key = p.get("sheet_key") or p["ref_sheet"]
        summary[key][bucket] += 1
        summary["ALL"][bucket] += 1
        func_summary[key][p["function"]][bucket] += 1
        func_summary[key][p["function"]]["total"] += 1
        summary[key]["total"] += 1
        summary["ALL"]["total"] += 1
        if bucket in {"full_inconsistent", "mixed"}:
            summary[key]["inconsistent"] += 1
            summary["ALL"]["inconsistent"] += 1
            func_summary[key][p["function"]]["inconsistent"] += 1
    return summary, func_summary


def sheet_keys_in_order(params):
    order = []
    seen = set()
    for p in params:
        key = p.get("sheet_key") or p["ref_sheet"]
        if key not in seen:
            seen.add(key)
            order.append(key)
    return order


def network_for_sheet(params, sheet_key):
    for p in params:
        if (p.get("sheet_key") or p["ref_sheet"]) == sheet_key:
            return p.get("network") or ""
    return ""


def files_banner(run: RunContext) -> str:
    refs = ", ".join(p.name for p in run.reference_files) or run.reference.name
    inputs = ", ".join(workbook_rel_label(p, run.input_folder) for p in run.input_files)
    if not inputs:
        names = [p.name for p in (run.cfg_4g, run.cfg_5g) if p]
        inputs = ", ".join(names)
    rats = ", ".join(run.selected_rats or smart.RAT_FOLDERS)
    return f"{APP_TITLE} | Reference: {refs} | Input: {inputs} | Networks: {rats} | Run {run.run_id}"


def fill_header(ws, titles, fill, font):
    for col, title in enumerate(titles, start=1):
        cell = ws.cell(1, col, title)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autosize(ws, max_width=42):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 12
        for cell in col[:80]:
            if cell.value:
                width = min(max_width, max(width, min(len(str(cell.value)) + 2, max_width)))
        ws.column_dimensions[letter].width = width


def status_fill(status):
    colors = {
        "Consistent": "C6EFCE",
        "Inconsistent": "FFC7CE",
        "Mixed / Partial": "FFEB9C",
        "Not Found in Configuration": "F4B183",
        "No Recommend Value": "D9D9D9",
    }
    return PatternFill("solid", fgColor=colors.get(status, "FFFFFF"))


def write_excel(params, summary, func_summary, run: RunContext):
    run.out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(bold=True, size=14, color="1F4E79")
    section_font = Font(bold=True, size=12, color="1F4E79")

    # 1. Overall Report
    ws = wb.active
    ws.title = "1_Overall_Report"
    ws["A1"] = "Parameter Inconsistency Overall Report"
    ws["A1"].font = title_font
    ws.merge_cells("A1:G1")
    ws["A2"] = files_banner(run)
    ws.merge_cells("A2:G2")
    if run.input_folder or run.reference_folder:
        ws["A3"] = (
            f"Input Folder: {run.input_folder or '(files)'} | "
            f"Reference Folder: {run.reference_folder or '(files)'} | "
            f"Output Folder: {run.output_dir}"
        )
        ws.merge_cells("A3:J3")

    ws["A4"] = "1.1 Overall Summary (every reference file / sheet vs every input file)"
    ws["A4"].font = section_font
    headers = [
        "Reference File / Sheet",
        "Network",
        "Total Parameters",
        "Inconsistent (total)",
        "Full Inconsistent",
        "Mixed / Partial",
        "Consistent",
        "Not Found in Config",
        "No Recommend Value",
        "Inconsistency Rate (auditable)",
    ]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(5, col, title)
        cell.fill = header_fill
        cell.font = header_font

    def auditable_rate(counter):
        auditable = counter["inconsistent"] + counter["consistent"]
        if not auditable:
            return "N/A"
        return f"{counter['inconsistent'] / auditable:.1%}"

    keys = sheet_keys_in_order(params)
    rows = [(key, network_for_sheet(params, key), summary[key]) for key in keys]
    rows.append(("ALL FILES / SHEETS", "ALL", summary["ALL"]))
    for idx, (name, net, counter) in enumerate(rows, start=6):
        values = [
            name,
            net,
            counter["total"],
            counter["inconsistent"],
            counter["full_inconsistent"],
            counter["mixed"],
            counter["consistent"],
            counter["not_found"],
            counter["no_recommend"],
            auditable_rate(counter),
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(idx, col, val)
            cell.border = thin
            if name == "ALL FILES / SHEETS":
                cell.font = Font(bold=True)
            if col == 4 and counter["inconsistent"]:
                cell.fill = PatternFill("solid", fgColor="FFC7CE")

    note_row = 6 + len(rows) + 1
    ws.cell(note_row, 1, "Notes").font = section_font
    ws.cell(
        note_row + 1,
        1,
        "Every workbook in the selected Input network folders (5G/4G/3G/2G) is compared against every workbook in the Reference Folder. "
        "Band-specific Proposed/Recommended values (L09/L18/L21/L26) are applied using Cell (4G) and NRDUCell (5G). "
        "Auditable parameters = parameters that have a Recommend Value and were found in configuration. "
        "Inconsistent (total) = Full Inconsistent + Mixed/Partial. "
        "Full Inconsistent = every object differs from recommend. "
        "Mixed/Partial = some cells/sites match and others do not. "
        "Parameters with empty Recommend Value are listed but not counted as inconsistency. "
        f"Excel sheet limit (Microsoft): {EXCEL_MAX_ROWS:,} rows x {EXCEL_MAX_COLS:,} columns. "
        "This tool does not add a lower file, sheet, or column cap.",
    )
    ws.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 2, end_column=10)
    ws.cell(note_row + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    autosize(ws)
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A6"

    # 2. All parameter wise
    ws2 = wb.create_sheet("2_All_Parameter_Report")
    ws2["A1"] = "All Parameter-wise Report (function wise)"
    ws2["A1"].font = title_font
    cond_slots = condition_slots_for_report(params)
    cond_headers = [f"Conditions{n}" for n in cond_slots]
    param_headers = [
        "Reference File",
        "Reference Sheet",
        "Network",
        "Function",
        "MML Object",
        "Parameter ID",
        "Parameter Name",
        "Recommend Value",
        *cond_headers,
        "Applied Recommend (band/group)",
        "Status",
        "Input File(s)",
        "Config Sheet",
        "Config Column",
        "Bit Name",
        "Objects Checked",
        "Match Count",
        "Mismatch Count",
        "Match %",
        "Actual Value Distribution (top)",
        "Sample Mismatched Objects",
        "Remark",
    ]
    status_col = 8 + len(cond_slots) + 2
    for col, title in enumerate(param_headers, start=1):
        cell = ws2.cell(3, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws2.row_dimensions[3].height = 32
    ordered = sorted(
        params,
        key=lambda p: (p.get("ref_file") or "", p["ref_sheet"], p["function"], p["mml"], p["pid"]),
    )
    r = 4
    for p in ordered:
        res = p["result"]
        resolved = p["resolved"]
        match_pct = ""
        if res["objects_checked"] and not is_empty(p["recommend"]):
            match_pct = f"{res['match_count'] / res['objects_checked']:.1%}"
        values = [
            p.get("ref_file") or run.reference.name,
            p["ref_sheet"],
            p.get("network") or "",
            p["function"],
            p["mml"],
            p["pid"],
            p.get("param_name") or "",
            "" if p["recommend"] is None else p["recommend"],
            *[condition_value_for_slot(p, n) or "" for n in cond_slots],
            res.get("applied_recommend") or "",
            res["status"],
            res.get("input_files") or "",
            resolved.get("sheet") or "",
            resolved.get("column") or "",
            resolved.get("bit") or "",
            res["objects_checked"],
            res["match_count"],
            res["mismatch_count"],
            match_pct,
            unique_actual_text(res["unique_actuals"]),
            " || ".join(res["sample_mismatches"]),
            res["remark"],
        ]
        for col, val in enumerate(values, start=1):
            cell = ws2.cell(r, col, val)
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col == status_col:
                cell.fill = status_fill(res["status"])
        r += 1
    last_col = get_column_letter(len(param_headers))
    ws2.freeze_panes = "A4"
    ws2.auto_filter.ref = f"A3:{last_col}{max(r-1, 3)}"
    autosize(ws2, 36)
    ws2.column_dimensions["F"].width = 42
    ws2.column_dimensions["G"].width = 36
    cond_last = 8 + len(cond_slots)
    if cond_slots:
        for idx in range(9, cond_last + 1):
            ws2.column_dimensions[get_column_letter(idx)].width = 36
    applied_col = get_column_letter(cond_last + 1)
    ws2.column_dimensions[applied_col].width = 36

    # 3. Sheet wise
    ws3 = wb.create_sheet("3_Sheet_Wise_Report")
    ws3["A1"] = "Sheet-wise Report"
    ws3["A1"].font = title_font
    row = 3
    for sheet_name in keys:
        network = network_for_sheet(params, sheet_name)
        counter = summary[sheet_name]
        title = f"{sheet_name}" + (f"  ({network})" if network else "")
        ws3.cell(row, 1, title).font = section_font
        row += 1
        ws3.cell(row, 1, f"Compared against all input files: {', '.join(p.name for p in run.input_files)}")
        row += 1
        stats = [
            ("Total reference parameters", counter["total"]),
            ("Inconsistent (incl. Mixed)", counter["inconsistent"]),
            ("Consistent", counter["consistent"]),
            ("Not Found in Configuration", counter["not_found"]),
            ("No Recommend Value", counter["no_recommend"]),
            ("Inconsistency rate of auditable parameters", auditable_rate(counter)),
        ]
        for label, val in stats:
            ws3.cell(row, 1, label)
            ws3.cell(row, 2, val)
            if label.startswith("Inconsistent"):
                ws3.cell(row, 2).fill = PatternFill("solid", fgColor="FFC7CE")
            row += 1
        row += 1
        ws3.cell(row, 1, "Function").font = header_font
        ws3.cell(row, 1).fill = header_fill
        for col, title in enumerate(
            ["Function", "Total", "Inconsistent", "Consistent", "Not Found", "No Recommend", "Rate"],
            start=1,
        ):
            cell = ws3.cell(row, col, title)
            cell.fill = header_fill
            cell.font = header_font
        row += 1
        for func in sorted(func_summary[sheet_name].keys()):
            c = func_summary[sheet_name][func]
            vals = [
                func,
                c["total"],
                c["inconsistent"],
                c["consistent"],
                c["not_found"],
                c["no_recommend"],
                auditable_rate(c),
            ]
            for col, val in enumerate(vals, start=1):
                cell = ws3.cell(row, col, val)
                cell.border = thin
            row += 1
        row += 1
        ws3.cell(row, 1, f"Inconsistent parameters in {sheet_name}").font = section_font
        row += 1
        for col, title in enumerate(
            ["Function", "MML Object", "Parameter ID", "Recommend", "Status", "Actual (top)", "Mismatch Count", "Input File(s)"],
            start=1,
        ):
            cell = ws3.cell(row, col, title)
            cell.fill = header_fill
            cell.font = header_font
        row += 1
        sheet_params = [
            p
            for p in ordered
            if (p.get("sheet_key") or p["ref_sheet"]) == sheet_name and is_inconsistent(p["result"]["status"])
        ]
        if not sheet_params:
            ws3.cell(row, 1, "None")
            row += 1
        for p in sheet_params:
            vals = [
                p["function"],
                p["mml"],
                p["pid"],
                "" if p["recommend"] is None else p["recommend"],
                p["result"]["status"],
                unique_actual_text(p["result"]["unique_actuals"]),
                p["result"]["mismatch_count"],
                p["result"].get("input_files") or "",
            ]
            for col, val in enumerate(vals, start=1):
                cell = ws3.cell(row, col, val)
                cell.border = thin
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if col == 5:
                    cell.fill = status_fill(p["result"]["status"])
            row += 1
        row += 2
    autosize(ws3, 40)
    ws3.column_dimensions["C"].width = 45
    ws3.column_dimensions["F"].width = 50

    # 4. Final Summary
    ws4 = wb.create_sheet("4_Final_Summary")
    ws4["A1"] = "Final Summary"
    ws4["A1"].font = title_font
    all_c = summary["ALL"]
    inconsistent_params = [p for p in ordered if is_inconsistent(p["result"]["status"])]
    consistent_params = [p for p in ordered if p["result"]["status"] == "Consistent"]
    not_found = [p for p in ordered if p["result"]["status"] == "Not Found in Configuration"]
    no_rec = [p for p in ordered if p["result"]["status"] == "No Recommend Value"]

    ref_names = ", ".join(p.name for p in run.reference_files) or run.reference.name
    in_names = ", ".join(p.name for p in run.input_files)
    lines = [
        f"Audit date: {run.audit_date}",
        f"Run ID: {run.run_id}",
        f"Reference Folder: {run.reference_folder or '(selected files)'}",
        f"Reference files ({len(run.reference_files)}): {ref_names}",
        f"Input Folder: {run.input_folder or '(selected files)'}",
        f"Input files ({len(run.input_files)}): {in_names}",
        f"Output Folder: {run.output_dir}",
        "",
        "HEADLINE",
        f"- Total reference parameters checked: {all_c['total']}",
        f"- Parameter inconsistencies (value mismatch or mixed): {all_c['inconsistent']}",
        f"  of which full inconsistent: {all_c['full_inconsistent']}, mixed/partial: {all_c['mixed']}",
        f"- Consistent parameters: {all_c['consistent']}",
        f"- Not found in configuration dumps: {all_c['not_found']}",
        f"- No recommend value (not audited): {all_c['no_recommend']}",
        f"- Combined auditable inconsistency rate: {auditable_rate(all_c)}",
        "",
        "BY REFERENCE SHEET",
    ]
    for sheet_name in keys:
        c = summary[sheet_name]
        net = network_for_sheet(params, sheet_name)
        net_txt = f" ({net})" if net else ""
        lines.append(
            f"- {sheet_name}{net_txt}: {c['total']} parameters; {c['inconsistent']} inconsistent; "
            f"{c['consistent']} consistent; {c['not_found']} not found; {c['no_recommend']} no recommend; "
            f"auditable rate {auditable_rate(c)}"
        )
    lines += [
        "",
        "FUNCTION HOTSPOTS",
    ]
    hotspots = []
    for sheet_name in keys:
        for func, c in func_summary[sheet_name].items():
            if c["inconsistent"]:
                hotspots.append((c["inconsistent"], sheet_name, func, c["total"]))
    hotspots.sort(reverse=True)
    for count, sheet_name, func, total in hotspots[:12]:
        lines.append(f"- {sheet_name} / {func}: {count} inconsistent of {total} parameters")

    lines += [
        "",
        "INTERPRETATION",
        "- Every reference workbook is fully analyzed. Every input workbook is searched for each parameter.",
        "- Sheet names in input dumps that match an MO / MML Object name are treated as that object.",
        "- Huawei dumps with a MAPPING DEF sheet are mapped by MOC/attribute; other workbooks use sheet and column names.",
        "- Band-conditional values (L9:-74 L18:-118 …) use Local cell ID looked up in Cell (4G) / NRDUCELL (5G) only. Other MO sheets are not preloaded for band mapping.",
        "- Switch bits are extracted from Huawei bit-pack strings (NAME-1&NAME-0). Recommend may name the bit as SwitchName-0 / SwitchName-1 (or several joined by &) against Parameter Name; Parameter ID BIT@Attribute still works.",
        "- Mixed / Partial means some cells or sites match the recommend value and others do not; these are counted as inconsistency.",
        "- Empty Recommend Value in the reference file cannot be judged; they are excluded from the inconsistency rate.",
        f"- Microsoft Excel limit: {EXCEL_MAX_ROWS:,} rows and {EXCEL_MAX_COLS:,} columns per sheet. This tool does not add a lower cap.",
        "",
        f"See sheet 2_All_Parameter_Report for the full {all_c['total']} parameter lines, "
        f"and sheet 3_Sheet_Wise_Report for the inconsistent list of each reference sheet.",
    ]
    for idx, line in enumerate(lines, start=3):
        ws4.cell(idx, 1, line)
        if line in {"HEADLINE", "BY REFERENCE SHEET", "FUNCTION HOTSPOTS", "INTERPRETATION"}:
            ws4.cell(idx, 1).font = section_font
    ws4.column_dimensions["A"].width = 140
    ws4["A3"].alignment = Alignment(wrap_text=False)

    # helper lists
    ws5 = wb.create_sheet("5_Not_Found_No_Recommend")
    ws5["A1"] = "Parameters not found or with no recommend value"
    ws5["A1"].font = title_font
    for col, title in enumerate(
        ["Category", "Reference Sheet", "Function", "MML Object", "Parameter ID", "Recommend", "Remark"],
        start=1,
    ):
        cell = ws5.cell(3, col, title)
        cell.fill = header_fill
        cell.font = header_font
    r = 4
    for p in not_found + no_rec:
        vals = [
            p["result"]["status"],
            p["ref_sheet"],
            p["function"],
            p["mml"],
            p["pid"],
            "" if p["recommend"] is None else p["recommend"],
            p["result"]["remark"],
        ]
        for col, val in enumerate(vals, start=1):
            cell = ws5.cell(r, col, val)
            cell.border = thin
            if col == 1:
                cell.fill = status_fill(p["result"]["status"])
        r += 1
    autosize(ws5, 40)

    wb.save(run.out_xlsx)
    return {
        "inconsistent_params": inconsistent_params,
        "consistent_params": consistent_params,
        "not_found": not_found,
        "no_rec": no_rec,
        "hotspots": hotspots,
    }


def write_markdown(params, summary, func_summary, extras, run: RunContext):
    def rate(counter):
        auditable = counter["inconsistent"] + counter["consistent"]
        if not auditable:
            return "N/A"
        return f"{counter['inconsistent'] / auditable:.1%}"

    all_c = summary["ALL"]
    keys = sheet_keys_in_order(params)
    ref_names = ", ".join(f"`{p.name}`" for p in run.reference_files) or f"`{run.reference.name}`"
    in_names = ", ".join(f"`{p.name}`" for p in run.input_files)
    lines = []
    lines.append(f"# Parameter Inconsistency Report ({run.audit_date})")
    lines.append("")
    lines.append(f"Run ID: `{run.run_id}`")
    lines.append(f"Reference files: {ref_names}")
    lines.append(f"Input files: {in_names}")
    if run.input_folder:
        lines.append(f"Input Folder: `{run.input_folder}`")
    if run.reference_folder:
        lines.append(f"Reference Folder: `{run.reference_folder}`")
    lines.append(f"Output Folder: `{run.output_dir}`")
    lines.append("")
    lines.append("## 1. Overall Report")
    lines.append("")
    lines.append("### 1.1 How many parameter inconsistencies from all reference sheets")
    lines.append("")
    lines.append("| Reference File / Sheet | Network | Total | Inconsistent | Full | Mixed | Consistent | Not Found | No Recommend | Auditable Inconsistency Rate |")
    lines.append("|---|---|---:|---:|---:|---:|---:|---:|---:|---|")
    for name in keys:
        net = network_for_sheet(params, name)
        c = summary[name]
        lines.append(
            f"| {name} | {net} | {c['total']} | {c['inconsistent']} | {c['full_inconsistent']} | {c['mixed']} | "
            f"{c['consistent']} | {c['not_found']} | {c['no_recommend']} | {rate(c)} |"
        )
    lines.append(
        f"| ALL FILES / SHEETS | ALL | {all_c['total']} | {all_c['inconsistent']} | {all_c['full_inconsistent']} | {all_c['mixed']} | "
        f"{all_c['consistent']} | {all_c['not_found']} | {all_c['no_recommend']} | {rate(all_c)} |"
    )
    lines.append("")
    lines.append(
        f"**Headline:** {all_c['inconsistent']} of {all_c['total']} reference parameters are inconsistent "
        f"({all_c['full_inconsistent']} full mismatch, {all_c['mixed']} mixed/partial). "
        f"{all_c['consistent']} match the recommend value on all objects. "
        f"{all_c['no_recommend']} have no recommend value and were not audited. "
        f"{all_c['not_found']} could not be located in the configuration dumps."
    )
    lines.append("")

    lines.append("## 2. All Parameter-wise Report (function wise)")
    lines.append("")
    lines.append(f"Full line-by-line table is in `{run.out_xlsx}` sheet `2_All_Parameter_Report`.")
    lines.append("Below: every parameter grouped by reference sheet and function.")
    lines.append("")
    ordered = sorted(params, key=lambda p: (p.get("ref_file") or "", p["ref_sheet"], p["function"], p["mml"], p["pid"]))
    current = None
    for p in ordered:
        key = (p.get("sheet_key") or p["ref_sheet"], p["function"])
        if key != current:
            current = key
            lines.append(f"### {p.get('sheet_key') or p['ref_sheet']} — {p['function']}")
            lines.append("")
            lines.append("| Parameter ID | MML Object | Recommend | Status | Match | Mismatch | Actual (top) |")
            lines.append("|---|---|---|---|---:|---:|---|")
        rec = "" if p["recommend"] is None else str(p["recommend"])
        actual = unique_actual_text(p["result"]["unique_actuals"]).replace("|", "/")
        lines.append(
            f"| `{p['pid']}` | {p['mml']} | {rec} | {p['result']['status']} | "
            f"{p['result']['match_count']} | {p['result']['mismatch_count']} | {actual} |"
        )
    lines.append("")

    lines.append("## 3. Sheet-wise Report")
    lines.append("")
    for sheet_name in keys:
        c = summary[sheet_name]
        network = network_for_sheet(params, sheet_name)
        heading = f"{sheet_name} ({network})" if network else sheet_name
        lines.append(f"### {heading}")
        lines.append("")
        lines.append(f"- Total parameters: **{c['total']}**")
        lines.append(f"- Inconsistent: **{c['inconsistent']}**")
        lines.append(f"- Consistent: **{c['consistent']}**")
        lines.append(f"- Not found: **{c['not_found']}**")
        lines.append(f"- No recommend value: **{c['no_recommend']}**")
        lines.append(f"- Auditable inconsistency rate: **{rate(c)}**")
        lines.append("")
        lines.append("Inconsistent parameters:")
        lines.append("")
        lines.append("| Function | Parameter ID | Recommend | Status | Mismatch Count | Actual (top) |")
        lines.append("|---|---|---|---|---:|---|")
        sheet_incon = [
            p
            for p in ordered
            if (p.get("sheet_key") or p["ref_sheet"]) == sheet_name and is_inconsistent(p["result"]["status"])
        ]
        if not sheet_incon:
            lines.append("| — | — | — | None | 0 | — |")
        for p in sheet_incon:
            rec = "" if p["recommend"] is None else str(p["recommend"])
            actual = unique_actual_text(p["result"]["unique_actuals"]).replace("|", "/")
            lines.append(
                f"| {p['function']} | `{p['pid']}` | {rec} | {p['result']['status']} | "
                f"{p['result']['mismatch_count']} | {actual} |"
            )
        lines.append("")

    lines.append("## 4. Final Summary")
    lines.append("")
    sheet_bits = []
    for sheet_name in keys:
        c = summary[sheet_name]
        sheet_bits.append(f"{sheet_name}: {c['inconsistent']} inconsistent of {c['total']}")
    lines.append(
        f"The reference files contain **{all_c['total']}** parameters. "
        f"**{all_c['inconsistent']}** parameters are inconsistent against the input configuration dumps "
        f"({all_c['full_inconsistent']} fully mismatched, {all_c['mixed']} mixed). "
        + (" ".join(sheet_bits) + ". " if sheet_bits else "")
        + f"**{all_c['consistent']}** parameters fully match the recommend value. "
        f"**{all_c['no_recommend']}** parameters have no recommend value in the reference and were excluded "
        f"from the inconsistency rate. **{all_c['not_found']}** parameters could not be mapped to the dumps."
    )
    lines.append("")
    if extras["hotspots"]:
        lines.append("Largest function-level inconsistency counts:")
        lines.append("")
        for count, sheet_name, func, total in extras["hotspots"][:10]:
            lines.append(f"- **{sheet_name} / {func}**: {count} of {total} parameters inconsistent")
        lines.append("")
    lines.append(f"Detailed workbook: `{run.out_xlsx}`")
    run.out_md.write_text("\n".join(lines), encoding="utf-8")


def write_parameter_csv(params, run: RunContext):
    cond_slots = condition_slots_for_report(params)
    cond_fields = [f"conditions{n}" for n in cond_slots]
    fieldnames = [
        "run_id",
        "reference_file",
        "reference_sheet",
        "network",
        "function",
        "mml_object",
        "parameter_id",
        "parameter_name",
        "recommend_value",
        *cond_fields,
        "applied_recommend",
        "status",
        "input_files",
        "config_sheet",
        "config_column",
        "bit_name",
        "objects_checked",
        "match_count",
        "mismatch_count",
        "actual_top",
        "remark",
    ]
    with run.out_csv.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for p in params:
            res = p["result"]
            resolved = p["resolved"]
            row = {
                "run_id": run.run_id,
                "reference_file": p.get("ref_file") or run.reference.name,
                "reference_sheet": p["ref_sheet"],
                "network": p.get("network") or "",
                "function": p["function"],
                "mml_object": p["mml"],
                "parameter_id": p["pid"],
                "parameter_name": p.get("param_name") or "",
                "recommend_value": "" if p["recommend"] is None else p["recommend"],
                "applied_recommend": res.get("applied_recommend") or "",
                "status": res["status"],
                "input_files": res.get("input_files") or "",
                "config_sheet": resolved.get("sheet") or "",
                "config_column": resolved.get("column") or "",
                "bit_name": resolved.get("bit") or "",
                "objects_checked": res["objects_checked"],
                "match_count": res["match_count"],
                "mismatch_count": res["mismatch_count"],
                "actual_top": unique_actual_text(res["unique_actuals"]),
                "remark": res["remark"],
            }
            for n in cond_slots:
                row[f"conditions{n}"] = condition_value_for_slot(p, n) or ""
            writer.writerow(row)


def rate_text(counter) -> str:
    auditable = counter["inconsistent"] + counter["consistent"]
    if not auditable:
        return "N/A"
    return f"{counter['inconsistent'] / auditable:.1%}"


def write_run_summary(summary, run: RunContext, param_count: int):
    payload = {
        "run_id": run.run_id,
        "audit_date": run.audit_date,
        "input_folder": str(run.input_folder) if run.input_folder else None,
        "reference_folder": str(run.reference_folder) if run.reference_folder else None,
        "output_folder": str(run.output_dir),
        "reference_files": [str(p) for p in run.reference_files],
        "input_files": [str(p) for p in run.input_files],
        "selected_rats": list(run.selected_rats),
        "app_version": APP_VERSION,
        "parameter_count": param_count,
        "overall": dict(summary["ALL"]),
        "sheets": {key: dict(counter) for key, counter in summary.items() if key != "ALL"},
        "auditable_inconsistency_rate": rate_text(summary["ALL"]),
        "excel_limits_note": f"{EXCEL_MAX_ROWS} rows x {EXCEL_MAX_COLS} columns per sheet (Microsoft Excel). No extra cap in this tool.",
        "outputs": {
            "xlsx": str(run.out_xlsx),
            "md": str(run.out_md),
            "csv": str(run.out_csv),
        },
    }
    run.out_summary.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return payload


def append_history(summary, run: RunContext):
    all_c = summary["ALL"]
    fieldnames = [
        "run_id",
        "audit_date",
        "reference_files",
        "input_files",
        "total",
        "inconsistent",
        "full_inconsistent",
        "mixed",
        "consistent",
        "not_found",
        "no_recommend",
        "auditable_rate",
        "report_xlsx",
    ]
    new_file = not run.history_csv.exists()
    with run.history_csv.open("a", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        if new_file:
            writer.writeheader()
        try:
            report_path = str(run.out_xlsx.relative_to(run.root))
        except ValueError:
            report_path = str(run.out_xlsx)
        writer.writerow(
            {
                "run_id": run.run_id,
                "audit_date": run.audit_date,
                "reference_files": "; ".join(p.name for p in run.reference_files) or run.reference.name,
                "input_files": "; ".join(p.name for p in run.input_files),
                "total": all_c["total"],
                "inconsistent": all_c["inconsistent"],
                "full_inconsistent": all_c["full_inconsistent"],
                "mixed": all_c["mixed"],
                "consistent": all_c["consistent"],
                "not_found": all_c["not_found"],
                "no_recommend": all_c["no_recommend"],
                "auditable_rate": rate_text(all_c),
                "report_xlsx": report_path,
            }
        )


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description="Compare every reference workbook with every input dump. "
        "Use three folders (Input, Reference, Output), or launch the GUI with no arguments."
    )
    parser.add_argument(
        "files",
        nargs="*",
        type=Path,
        help="Optional workbooks (legacy). Prefer --input-folder and --reference-folder.",
    )
    parser.add_argument("--input-folder", "--input-dir", dest="input_folder", type=Path, default=None, help="Folder of configuration dumps (any names)")
    parser.add_argument("--reference-folder", dest="reference_folder", type=Path, default=None, help="Folder of reference / plan-value workbooks (any names)")
    parser.add_argument("-o", "--output-dir", "--output-folder", dest="output_dir", type=Path, default=DEFAULT_OUTPUT_DIR, help="Output folder for reports")
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG, help="Optional JSON config with file patterns")
    parser.add_argument("--reference", type=Path, help="Single reference workbook (legacy)")
    parser.add_argument("--config-4g", type=Path, help="Single 4G dump (legacy)")
    parser.add_argument("--config-5g", type=Path, help="Single 5G dump (legacy)")
    parser.add_argument("--run-id", help="Optional run id; default is timestamp YYYYMMDD_HHMMSS")
    parser.add_argument(
        "--rats",
        default="5G,4G,3G,2G",
        help="Which Input subfolders to search: 5G,4G,3G,2G (comma-separated). Default: all.",
    )
    parser.add_argument("--license", type=Path, help="ParameterAudit.lic file (default: next to the app)")
    parser.add_argument("--license-status", action="store_true", help="Print license status and exit")
    parser.add_argument("--list-inputs", action="store_true", help="Show detected input files and exit")
    parser.add_argument("--gui", action="store_true", help="Open the graphical tool")
    parser.add_argument(
        "--fail-on-inconsistent",
        action="store_true",
        help="Exit with code 1 if any auditable parameter is inconsistent",
    )
    return parser.parse_args(argv)


def resolve_required_file(label: str, explicit: Path | None, discovered: Path | None) -> Path:
    path = explicit or discovered
    if path is None or not path.exists():
        raise FileNotFoundError(
            f"{label} not found. Put the file in input/ or the repo root, or pass an explicit path."
        )
    return path.resolve()


def execute_folder_audit(
    output_dir: Path,
    input_folder: Path | None = None,
    reference_folder: Path | None = None,
    input_files=None,
    reference_files=None,
    run_id=None,
    progress=print,
    license_file=None,
    rats=None,
):
    license_info = license_mod.require_active_license(license_file)
    progress(f"{APP_TITLE}")
    progress(f"License: {license_info.message}")
    selected_rats = smart.normalize_rats(rats)
    if not selected_rats:
        raise ValueError("Select at least one network: 5G, 4G, 3G, or 2G.")
    if input_folder:
        ensure_rat_input_folders(input_folder)
    input_files = [Path(p).expanduser().resolve() for p in (input_files or [])]
    reference_files = [Path(p).expanduser().resolve() for p in (reference_files or [])]
    if not input_files:
        input_files = list_workbooks(input_folder, rats=selected_rats, use_rat_subfolders=True)
    if not reference_files:
        reference_files = list_workbooks(reference_folder)
    if not input_files:
        searched = ", ".join(f"{input_folder}/{rat}" if input_folder else rat for rat in selected_rats)
        raise FileNotFoundError(
            "No configuration dumps found in the selected network folders "
            f"({', '.join(selected_rats)}). Put 5G dumps in Input/5G, 4G dumps in Input/4G, "
            f"3G in Input/3G, 2G in Input/2G. Searched: {searched}."
        )
    if not reference_files:
        raise FileNotFoundError(
            "Reference Folder has no Excel workbooks (.xlsx / .xlsb / .xlsm). "
            "Put every recommended/plan-value file there."
        )

    guessed_4g = next((p for p in input_files if guess_file_role(p) == "4g"), input_files[0])
    guessed_5g = next((p for p in input_files if guess_file_role(p) == "5g"), input_files[-1])
    run = build_run_context(
        root=ROOT,
        reference=reference_files[0],
        cfg_4g=guessed_4g,
        cfg_5g=guessed_5g,
        output_dir=Path(output_dir),
        run_id=run_id,
        input_folder=input_folder,
        reference_folder=reference_folder,
        input_files=input_files,
        reference_files=reference_files,
        selected_rats=selected_rats,
    )
    progress(f"Run ID: {run.run_id}")
    progress(f"Networks: {', '.join(selected_rats)}")
    progress(f"Reference Folder: {run.reference_folder or '(files)'}")
    for path in reference_files:
        progress(f"  REF  {path.name}")
    progress(f"Input Folder: {run.input_folder or '(files)'}")
    for path in input_files:
        progress(f"  IN   {workbook_rel_label(path, input_folder)}")
    progress(f"Output Folder: {run.output_dir}")
    progress("Loading reference parameters from every reference workbook...")
    params, missing_ref_sheets = load_all_references(reference_files)
    if missing_ref_sheets:
        progress(
            "  No parameter rows detected in: "
            + ", ".join(missing_ref_sheets)
            + " (need Parameter ID + Recommend/Proposed Value or MML Object columns)"
        )
    progress(f"  {len(params)} parameters from {len(reference_files)} reference file(s)")
    if not params:
        raise ValueError(
            "No reference parameters found. Each reference sheet needs headers such as "
            "Function / MML Object / Parameter ID / Parameter Name, with Proposed or Recommend Value found adaptively."
        )
    progress("Mapping Local cell ID to L09/L18/L21/L26 from Cell (4G) and NRDUCELL (5G) only...")
    store = InputStore(input_files)
    for wb in store.workbooks:
        band_map = wb.cell_index()
        progress(
            f"  {workbook_rel_label(wb.path, input_folder)}: "
            f"{len(band_map.by_cell_id)} cells from {', '.join(band_map._loaded_sheets) or 'no Cell/NRDUCELL sheet'}"
        )
    progress("Comparing each reference parameter against selected input files...")
    total = len(params)
    for i, param in enumerate(params, start=1):
        resolved, result = store.compare_param(param)
        param["resolved"] = resolved
        param["result"] = result
        label = param.get("param_name") or param.get("pid") or param.get("mml") or ""
        progress(f"  {i}/{total}  {label}")
    summary, func_summary = summarize(params)
    progress("Writing reports to the Output Folder...")
    extras = write_excel(params, summary, func_summary, run)
    write_markdown(params, summary, func_summary, extras, run)
    write_parameter_csv(params, run)
    write_run_summary(summary, run, len(params))
    append_history(summary, run)
    copy_latest(run)
    progress("DONE")
    progress(f"Excel: {run.out_xlsx}")
    progress(f"Markdown: {run.out_md}")
    progress(f"CSV: {run.out_csv}")
    progress(f"Latest copy: {run.latest_dir}")
    progress("ALL " + str(dict(summary["ALL"])))
    return run, summary, extras


def execute_audit(reference: Path, cfg_4g: Path, cfg_5g: Path, output_dir: Path, run_id=None, progress=print):
    return execute_folder_audit(
        output_dir=output_dir,
        input_files=[cfg_4g, cfg_5g],
        reference_files=[reference],
        run_id=run_id,
        progress=progress,
    )


def main(argv=None):
    args = parse_args(argv)
    if args.license_status:
        info = license_mod.license_status(args.license)
        print(info.message)
        if info.path:
            print(f"File: {info.path}")
        if info.active:
            print(f"Expires: {info.expires_at.isoformat()}")
        return 0 if info.active else 3
    try:
        license_mod.require_active_license(args.license)
    except license_mod.LicenseError as exc:
        print(exc, file=sys.stderr)
        print("Place ParameterAudit.lic next to the app, or pass --license. Ask the issuer to extend an expired license.", file=sys.stderr)
        return 3

    cfg = load_audit_config(args.config)
    input_folder = args.input_folder or DEFAULT_INPUT_DIR
    reference_folder = args.reference_folder or DEFAULT_REFERENCE_DIR
    discovered = discover_inputs(ROOT, input_folder, cfg)
    from_files = classify_input_files(args.files) if args.files else None

    if args.list_inputs:
        print("Input Folder:", input_folder)
        print("Networks:", args.rats)
        for path in list_workbooks(input_folder, rats=args.rats, use_rat_subfolders=True):
            print(f"  IN   {workbook_rel_label(path, input_folder)}")
        print("Reference Folder:", reference_folder)
        for path in list_workbooks(reference_folder):
            print(f"  REF  {path}")
        print("Search folders (legacy detect):")
        for folder in discovered["search_dirs"]:
            print(f"  {folder}")
        if from_files:
            print("From selected files:")
            print(f"  Reference: {from_files['reference']}")
            print(f"  4G config: {from_files['cfg_4g']}")
            print(f"  5G config: {from_files['cfg_5g']}")
            if from_files["unknown"]:
                print(f"  Unclassified: {from_files['unknown']}")
        print(f"Legacy Reference: {discovered['reference']}")
        print(f"Legacy 4G config: {discovered['cfg_4g']}")
        print(f"Legacy 5G config: {discovered['cfg_5g']}")
        return 0

    folder_inputs = list_workbooks(input_folder, rats=args.rats, use_rat_subfolders=True)
    folder_refs = list_workbooks(reference_folder)
    use_folders = bool(args.input_folder or args.reference_folder or (folder_inputs and folder_refs))

    try:
        if use_folders and folder_inputs and folder_refs and not args.files and not args.reference:
            run, summary, extras = execute_folder_audit(
                output_dir=args.output_dir,
                input_folder=input_folder,
                reference_folder=reference_folder,
                run_id=args.run_id,
                license_file=args.license,
                rats=args.rats,
            )
        else:
            extra_inputs = []
            extra_refs = []
            if from_files:
                extra_refs.extend(from_files.get("all_reference") or [])
                extra_inputs.extend(from_files.get("all_4g") or [])
                extra_inputs.extend(from_files.get("all_5g") or [])
                extra_inputs.extend(from_files.get("unknown") or [])
            reference = args.reference
            if reference:
                extra_refs.append(Path(reference))
            if args.config_4g:
                extra_inputs.append(Path(args.config_4g))
            if args.config_5g:
                extra_inputs.append(Path(args.config_5g))
            if not extra_refs:
                extra_refs = folder_refs or ([discovered["reference"]] if discovered["reference"] else [])
            if not extra_inputs:
                extra_inputs = folder_inputs
                if not extra_inputs:
                    extra_inputs = [p for p in (discovered["cfg_4g"], discovered["cfg_5g"]) if p]
            extra_refs = unique_files(extra_refs)
            extra_inputs = unique_files(extra_inputs)
            run, summary, extras = execute_folder_audit(
                output_dir=args.output_dir,
                input_folder=args.input_folder,
                reference_folder=args.reference_folder,
                input_files=extra_inputs,
                reference_files=extra_refs,
                run_id=args.run_id,
                license_file=args.license,
                rats=args.rats,
            )
    except (FileNotFoundError, ValueError, license_mod.LicenseError) as exc:
        print(exc, file=sys.stderr)
        return 2

    print(f"History: {run.history_csv}")
    for key, counter in summary.items():
        if key != "ALL":
            print(key, dict(counter))
    if extras["not_found"]:
        print("Not found:")
        for p in extras["not_found"]:
            print(" ", p.get("ref_file"), p["ref_sheet"], p["mml"], p["pid"], p["result"]["remark"])
    if args.fail_on_inconsistent and summary["ALL"]["inconsistent"]:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
